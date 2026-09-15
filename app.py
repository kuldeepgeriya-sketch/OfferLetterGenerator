from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, session
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from docx import Document
from num2words import num2words
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import sqlite3
import os
import io
import re
import shutil
from datetime import datetime, date, timedelta

app = Flask(__name__)
app.secret_key = os.environ.get('FLASK_SECRET_KEY', 'dev-only-key-do-not-use-in-prod')
ADMIN_USER = os.environ.get('ADMIN_USER', 'admin')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'admin')  # override this in production env vars
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "hrms.db")
TEMPLATE_PATH = os.path.join(BASE_DIR, "_Offer_Letter_BSR.docx")
LEAVE_TYPES = ["Earned Leave", "Special Leave", "Loss of Pay"]
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def get_connection():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def get_next_employee_code():
    """Generate next employee code based on the last one (e.g. EE/69 -> EE/70).
    Scans BOTH employees table AND employee_joining_forms (pending + any others)
    so that multiple new joiners submitting before any approvals still get
    unique sequential codes in the series (prevents EE/70 being given to two people).
    """
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT employee_code FROM employees")
    emp_codes = cursor.fetchall()
    cursor.execute("SELECT employee_code FROM employee_joining_forms WHERE employee_code IS NOT NULL")
    join_codes = cursor.fetchall()
    conn.close()
    
    max_num = 69  # base on last given EE/69
    for row in emp_codes + join_codes:
        code = row[0] if row[0] else ''
        if code and code.startswith('EE/'):
            try:
                num = int(code.split('/')[-1])
                if num > max_num:
                    max_num = num
            except:
                pass
    return f"EE/{max_num + 1}"


def ensure_column(cursor, table, column, definition):
    cursor.execute(f"PRAGMA table_info({table})")
    existing = [row["name"] for row in cursor.fetchall()]
    if column not in existing:
        cursor.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS departments ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT UNIQUE NOT NULL,"
        "description TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS offices ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT UNIQUE NOT NULL,"
        "location TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS designations ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "title TEXT NOT NULL,"
        "department_id INTEGER"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS shifts ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT NOT NULL,"
        "start_time TEXT,"
        "end_time TEXT,"
        "description TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS employees ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_code TEXT UNIQUE NOT NULL,"
        "first_name TEXT NOT NULL,"
        "last_name TEXT NOT NULL,"
        "email TEXT,"
        "phone TEXT,"
        "dob TEXT,"
        "gender TEXT,"
        "department_id INTEGER,"
        "designation_id INTEGER,"
        "office_id INTEGER,"
        "manager_id INTEGER,"
        "joining_date TEXT,"
        "ctc REAL,"
        "status TEXT,"
        "address TEXT,"
        "pf_active INTEGER DEFAULT 1,"
        "esic_active INTEGER DEFAULT 1,"
        "professional_tax_active INTEGER DEFAULT 1,"
        "lwf_active INTEGER DEFAULT 1"
        ")"
    )
    ensure_column(cursor, "employees", "pf_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "esic_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "professional_tax_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "lwf_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "password", "TEXT")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS employee_joining_forms ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "full_name TEXT,"
        "contact_no TEXT,"
        "email_id TEXT,"
        "designation TEXT,"
        "joining_date TEXT,"
        "permanent_address TEXT,"
        "date_of_birth TEXT,"
        "gender TEXT,"
        "marital_status TEXT,"
        "pan_no TEXT,"
        "aadhar_no TEXT,"
        "bank_name TEXT,"
        "bank_account_number TEXT,"
        "ifsc_code TEXT,"
        "branch_name TEXT,"
        "highest_qualification TEXT,"
        "university_board TEXT,"
        "year_of_passing TEXT,"
        "percentage_grade TEXT,"
        "emergency_person_name TEXT,"
        "emergency_person_mobile TEXT,"
        "emergency_relation TEXT,"
        "emergency_person_address TEXT,"
        "prev_company_name TEXT,"
        "prev_designation TEXT,"
        "prev_duration TEXT,"
        "prev_last_salary TEXT,"
        "photo_filename TEXT,"
        "pan_card_filename TEXT,"
        "aadhar_card_filename TEXT,"
        "cheque_passbook_filename TEXT,"
        "highest_education_cert_filename TEXT,"
        "last_3_month_salary_slip_filename TEXT,"
        "prev_employment_docs_filename TEXT,"
        "submitted_on TEXT"
        ")"
    )

    # For existing DBs with old schema, add new columns
    new_joining_cols = [
        ("highest_qualification", "TEXT"),
        ("university_board", "TEXT"),
        ("year_of_passing", "TEXT"),
        ("percentage_grade", "TEXT"),
        ("emergency_person_name", "TEXT"),
        ("emergency_person_mobile", "TEXT"),
        ("emergency_relation", "TEXT"),
        ("emergency_person_address", "TEXT"),
        ("prev_company_name", "TEXT"),
        ("prev_designation", "TEXT"),
        ("prev_duration", "TEXT"),
        ("prev_last_salary", "TEXT"),
        ("photo_filename", "TEXT"),
        ("pan_card_filename", "TEXT"),
        ("aadhar_card_filename", "TEXT"),
        ("cheque_passbook_filename", "TEXT"),
        ("highest_education_cert_filename", "TEXT"),
        ("last_3_month_salary_slip_filename", "TEXT"),
        ("prev_employment_docs_filename", "TEXT"),
    ]
    for col, defn in new_joining_cols:
        ensure_column(cursor, "employee_joining_forms", col, defn)

    ensure_column(cursor, "employee_joining_forms", "employee_code", "TEXT")
    ensure_column(cursor, "employee_joining_forms", "status", "TEXT DEFAULT 'pending'")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS attendance ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "date TEXT NOT NULL,"
        "status TEXT NOT NULL,"
        "shift_id INTEGER,"
        "check_in TEXT,"
        "check_out TEXT"
        ")"
    )
    ensure_column(cursor, "attendance", "remarks", "TEXT")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS leave_balances ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "earned_leave REAL DEFAULT 0,"
        "special_leave REAL DEFAULT 0,"
        "loss_of_pay REAL DEFAULT 0,"
        "year INTEGER NOT NULL,"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    ensure_column(cursor, "leave_balances", "special_leave", "REAL DEFAULT 0")
    # Migrate data from old casual_leave column if it exists (for backward compat)
    try:
        cursor.execute("PRAGMA table_info(leave_balances)")
        cols = [row["name"] for row in cursor.fetchall()]
        if "casual_leave" in cols:
            cursor.execute("UPDATE leave_balances SET special_leave = COALESCE(special_leave, 0) + COALESCE(casual_leave, 0)")
    except Exception:
        pass

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS leave_requests ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "leave_type TEXT NOT NULL,"
        "start_date TEXT NOT NULL,"
        "end_date TEXT NOT NULL,"
        "status TEXT DEFAULT 'Pending',"
        "reason TEXT,"
        "applied_on TEXT,"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS earned_leave_credits ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "year INTEGER NOT NULL,"
        "month INTEGER NOT NULL,"
        "credited REAL NOT NULL,"
        "credited_on TEXT,"
        "UNIQUE(employee_id, year, month),"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS special_leave_credits ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "year INTEGER NOT NULL,"
        "credited REAL NOT NULL,"
        "credited_on TEXT,"
        "UNIQUE(employee_id, year),"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS payrolls ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "month INTEGER NOT NULL,"
        "year INTEGER NOT NULL,"
        "ctc REAL NOT NULL,"
        "basic REAL NOT NULL,"
        "da REAL NOT NULL,"
        "hra REAL NOT NULL,"
        "conveyance REAL NOT NULL,"
        "special_allowance REAL NOT NULL,"
        "gross REAL NOT NULL,"
        "pf REAL NOT NULL,"
        "esic REAL NOT NULL,"
        "professional_tax REAL NOT NULL,"
        "lwf REAL NOT NULL,"
        "total_deductions REAL NOT NULL,"
        "net_pay REAL NOT NULL"
        ")"
    )

    default_departments = [
        ("Human Resources", "People operations and employee success."),
        ("Engineering", "Product and software development."),
        ("Finance", "Payroll, accounts and budgeting."),
        ("Sales", "Revenue generation and customer success.")
    ]
    for name, description in default_departments:
        try:
            cursor.execute("INSERT INTO departments (name, description) VALUES (?, ?)", (name, description))
        except sqlite3.IntegrityError:
            pass

    default_offices = [
        ("Head Office", "Downtown Campus"),
        ("Regional Office", "Sector 21")
    ]
    for name, location in default_offices:
        try:
            cursor.execute("INSERT INTO offices (name, location) VALUES (?, ?)", (name, location))
        except sqlite3.IntegrityError:
            pass

    default_designations = [
        ("HR Manager", 1),
        ("Software Engineer", 2),
        ("Accountant", 3),
        ("Sales Executive", 4)
    ]
    for title, dept_id in default_designations:
        cursor.execute(
            "SELECT id FROM designations WHERE title = ? AND department_id = ?",
            (title, dept_id)
        )
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO designations (title, department_id) VALUES (?, ?)",
                (title, dept_id)
            )

    default_shifts = [
        ("Day Shift", "09:00", "18:00", "Standard daytime shift."),
        ("Night Shift", "22:00", "06:00", "Overnight coverage shift.")
    ]
    for name, start_time, end_time, description in default_shifts:
        cursor.execute(
            "SELECT id FROM shifts WHERE name = ?",
            (name,)
        )
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO shifts (name, start_time, end_time, description) VALUES (?, ?, ?, ?)",
                (name, start_time, end_time, description)
            )

    # ---- Bank Bulk Salary File Generator tables (standalone employee list) ----
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS bank_employees ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT NOT NULL,"
        "ifsc TEXT NOT NULL,"
        "account TEXT NOT NULL,"
        "company TEXT NOT NULL,"
        "UNIQUE(name, company)"
        ")"
    )
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS bank_batches ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "batch_date TEXT NOT NULL,"
        "company TEXT NOT NULL,"
        "created_at TEXT NOT NULL,"
        "filename TEXT"
        ")"
    )
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS bank_transactions ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "batch_id INTEGER NOT NULL REFERENCES bank_batches(id) ON DELETE CASCADE,"
        "name TEXT NOT NULL,"
        "ifsc TEXT NOT NULL,"
        "account TEXT NOT NULL,"
        "amount REAL NOT NULL,"
        "mode TEXT NOT NULL,"
        "remarks TEXT NOT NULL"
        ")"
    )
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS bank_remarks ("
        "remark TEXT PRIMARY KEY,"
        "last_used TEXT NOT NULL,"
        "use_count INTEGER NOT NULL DEFAULT 1"
        ")"
    )

    # ---- Payroll Processing System tables (standalone employee master) ----
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS payroll_master_employees ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "code TEXT,"
        "name TEXT NOT NULL UNIQUE,"
        "office TEXT,"
        "ctc REAL NOT NULL"
        ")"
    )
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS payroll_runs ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "month TEXT NOT NULL,"
        "year TEXT NOT NULL,"
        "batch_no TEXT,"
        "saved_on TEXT NOT NULL"
        ")"
    )
    # Migrate an older payroll_runs schema (UNIQUE(month, year), no batch_no) if present —
    # that constraint blocked saving more than one batch per month, which is the bug we're fixing.
    cursor.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='payroll_runs'")
    _existing_sql_row = cursor.fetchone()
    if _existing_sql_row and _existing_sql_row[0] and "UNIQUE(month, year)" in _existing_sql_row[0]:
        cursor.execute("ALTER TABLE payroll_runs RENAME TO payroll_runs_old")
        cursor.execute(
            "CREATE TABLE payroll_runs ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT,"
            "month TEXT NOT NULL,"
            "year TEXT NOT NULL,"
            "batch_no TEXT,"
            "saved_on TEXT NOT NULL"
            ")"
        )
        cursor.execute(
            "INSERT INTO payroll_runs (id, month, year, saved_on) "
            "SELECT id, month, year, saved_on FROM payroll_runs_old"
        )
        cursor.execute("DROP TABLE payroll_runs_old")
    ensure_column(cursor, "payroll_runs", "batch_no", "TEXT")
    cursor.execute(
        "CREATE TABLE IF NOT EXISTS payroll_run_records ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "run_id INTEGER NOT NULL REFERENCES payroll_runs(id) ON DELETE CASCADE,"
        "name TEXT NOT NULL,"
        "code TEXT,"
        "office TEXT,"
        "ctc REAL,"
        "month_days INTEGER,"
        "present_days INTEGER,"
        "gross REAL,"
        "basic_da REAL,"
        "hra REAL,"
        "other_allowance REAL,"
        "total_earning REAL,"
        "employee_pf REAL,"
        "employee_esic REAL,"
        "pt REAL,"
        "total_deduction REAL,"
        "net_pay REAL,"
        "employer_pf REAL,"
        "employer_esic REAL,"
        "lwf REAL,"
        "cost_to_company REAL,"
        "additional_ctc REAL"
        ")"
    )

    # Backfill batch_no for any payroll_runs rows that predate this feature (e.g. migrated
    # rows above), numbering them in the order they were originally saved (by id).
    cursor.execute("SELECT id, month, year FROM payroll_runs WHERE batch_no IS NULL OR batch_no = '' ORDER BY id ASC")
    _rows_needing_batch_no = cursor.fetchall()
    if _rows_needing_batch_no:
        _month_abbrev = {
            "January": "Jan", "February": "Feb", "March": "Mar", "April": "Apr",
            "May": "May", "June": "Jun", "July": "Jul", "August": "Aug",
            "September": "Sep", "October": "Oct", "November": "Nov", "December": "Dec",
        }
        _serial_counters = {}
        for _row in _rows_needing_batch_no:
            _key = (_row["month"], _row["year"])
            _serial_counters[_key] = _serial_counters.get(_key, 0) + 1
            _abbr = _month_abbrev.get(_row["month"], str(_row["month"])[:3])
            _yy = str(_row["year"])[-2:]
            _batch_no = f"{_abbr}-{_yy}/{str(_serial_counters[_key]).zfill(2)}"
            cursor.execute("UPDATE payroll_runs SET batch_no = ? WHERE id = ?", (_batch_no, _row["id"]))

    conn.commit()
    conn.close()


def ensure_template():
    if os.path.exists(TEMPLATE_PATH):
        return

    doc = Document()
    doc.add_paragraph("Date: [DATE]")
    doc.add_paragraph("")
    doc.add_paragraph("Dear [PREFIX] [NAME],")
    doc.add_paragraph("")
    doc.add_paragraph("We are pleased to offer you a position as [POSITION] at our company.")
    doc.add_paragraph("")
    doc.add_paragraph("Position Details:")
    doc.add_paragraph("Joining Date: [JOINING_DATE]")
    doc.add_paragraph("Location: [LOCATION]")
    doc.add_paragraph("Department: [DEPARTMENT]")
    doc.add_paragraph("Monthly CTC: [CTC_M] ([CTC_M_WORD])")
    doc.add_paragraph("Annual CTC: [CTC_A] ([CTC_A_WORD])")
    doc.add_paragraph("")
    doc.add_paragraph("We look forward to welcoming you to the team.")
    doc.add_paragraph("")
    doc.add_paragraph("Best regards,")
    doc.add_paragraph("HR Team")
    doc.save(TEMPLATE_PATH)


def replace_placeholders(doc, replacements, bold_keys=None):
    if bold_keys is None:
        bold_keys = []

    def process_paragraph(paragraph):
        full_text = "".join(run.text for run in paragraph.runs)
        if any(key in full_text for key in replacements):
            for key, value in replacements.items():
                full_text = full_text.replace(key, value)
            for run in list(paragraph.runs):
                paragraph._element.remove(run._element)

            bold_ranges = []
            for bold_key in bold_keys:
                value = replacements.get(bold_key, "")
                if value and value in full_text:
                    start = full_text.find(value)
                    if start != -1:
                        bold_ranges.append((start, start + len(value)))
            bold_ranges.sort()

            cursor = 0
            for start, end in bold_ranges:
                if start > cursor:
                    paragraph.add_run(full_text[cursor:start])
                paragraph.add_run(full_text[start:end]).bold = True
                cursor = end
            if cursor < len(full_text):
                paragraph.add_run(full_text[cursor:])

    for paragraph in doc.paragraphs:
        process_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)
    for section in doc.sections:
        for header in section.header.paragraphs:
            process_paragraph(header)
        for footer in section.footer.paragraphs:
            process_paragraph(footer)


def dict_rows(rows):
    return [dict(row) for row in rows]


# ============================================================================
# ================  Bank Bulk Salary File Generator  =======================
# ============================================================================

BANK_COMPANIES = ["EE-Academy", "EE-Placement"]

BANK_COMPANY_CONFIG = {
    "EE-Academy":   {"code": "EXCD26", "account": "6055752772"},
    "EE-Placement": {"code": "EXCPLA", "account": "2749426820"},
}

BANK_COL_B_STATIC = "RPAY"
BANK_COL_I_STATIC = "M"
BANK_IFT_PREFIX = "KKBK"
BANK_DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
BANK_DEFAULT_REMARK = "Incentive"

_THIN = Side(style="thin", color="FF000000")
_MEDIUM = Side(style="medium", color="FF000000")
BANK_ROW_HEIGHT = 19.5

BANK_COLUMN_STYLES = {
    1: dict(font=Font(name="Arial", size=10, color="FF252524"),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(vertical="center", wrap_text=True),
            number_format="General"),
    2: dict(font=Font(name="Times New Roman", size=10),
            border=Border(left=None, right=_MEDIUM, top=None, bottom=_MEDIUM),
            alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
            number_format="General"),
    3: dict(font=Font(name="Calibri", size=10, color="FF000000"),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(horizontal="left", vertical="top"),
            number_format="General"),
    4: dict(font=Font(name="Calibri", size=14),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(vertical="center"),
            number_format="General"),
    5: dict(font=Font(name="Calibri", size=10),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(horizontal="left", vertical="top"),
            number_format="@"),
    6: dict(font=Font(name="Calibri", size=14),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(vertical="center"),
            number_format="General"),
    7: dict(font=Font(name="Arial", size=10, color="FF252524"),
            border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
            alignment=Alignment(vertical="center", wrap_text=True),
            number_format="General"),
    8: dict(font=Font(name="Arial", size=10, bold=True),
            border=Border(left=_MEDIUM, right=_MEDIUM, top=_MEDIUM, bottom=_MEDIUM),
            alignment=Alignment(wrap_text=True),
            number_format="General"),
    9: dict(font=Font(name="Calibri", size=10),
            border=Border(left=_THIN, right=None, top=_THIN, bottom=_THIN),
            alignment=Alignment(horizontal="left", vertical="top"),
            number_format="General"),
    10: dict(font=Font(name="Calibri", size=11),
             border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
             alignment=Alignment(),
             number_format="General"),
    11: dict(font=Font(name="Calibri", size=14),
             border=Border(left=None, right=_THIN, top=_THIN, bottom=_THIN),
             alignment=Alignment(vertical="center"),
             number_format="General"),
    12: dict(font=Font(name="Calibri", size=14),
             border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
             alignment=Alignment(vertical="center"),
             number_format="General"),
    13: dict(font=Font(name="Calibri", size=14),
             border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
             alignment=Alignment(vertical="center"),
             number_format="General"),
    14: dict(font=Font(name="Calibri", size=14),
             border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
             alignment=Alignment(vertical="center"),
             number_format="General"),
}
# Columns 15-23 (O..W) are blank spacer columns; give them a plain default style.
for _i in range(15, 24):
    BANK_COLUMN_STYLES[_i] = dict(
        font=Font(name="Calibri", size=10),
        border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
        alignment=Alignment(),
        number_format="General"
    )
BANK_COLUMN_STYLES[24] = dict(
    font=Font(name="Calibri", size=10),
    border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
    alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
    number_format="General"
)
BANK_COLUMN_STYLES[25] = dict(
    font=Font(name="Calibri", size=10),
    border=Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
    alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
    number_format="General"
)

BANK_COLUMN_WIDTHS = {
    "EE-Academy": {"A": 10, "B": 8, "C": 8, "E": 12, "G": 14, "H": 12, "K": 24, "M": 12, "N": 18, "X": 16, "Y": 16},
    "EE-Placement": {"A": 10, "B": 8, "C": 8, "E": 12, "G": 14, "H": 12, "K": 24, "M": 12, "N": 18, "X": 16, "Y": 16},
}


def bank_payment_mode_for_ifsc(ifsc):
    return "IFT" if str(ifsc).strip().upper().startswith(BANK_IFT_PREFIX) else "NEFT"


def bank_build_row(company, date_str, name, ifsc, account, amount, remarks):
    cfg = BANK_COMPANY_CONFIG[company]
    mode = bank_payment_mode_for_ifsc(ifsc)
    account_val = int(account) if str(account).isdigit() else str(account)
    row = [None] * 25
    row[0] = cfg["code"]
    row[1] = BANK_COL_B_STATIC
    row[2] = mode
    row[4] = date_str
    row[6] = int(cfg["account"])
    row[7] = amount
    row[8] = BANK_COL_I_STATIC
    row[10] = name
    row[12] = str(ifsc).strip().upper()
    row[13] = account_val
    row[23] = remarks
    row[24] = remarks
    return row


def bank_order_neft_then_ift(entries):
    return sorted(entries, key=lambda e: 0 if bank_payment_mode_for_ifsc(e["ifsc"]) == "NEFT" else 1)


def generate_bank_excel(company, date_str, entries):
    """Returns a BytesIO of the generated bulk salary Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ordered = bank_order_neft_then_ift(entries)

    for i, e in enumerate(ordered, start=1):
        row_vals = bank_build_row(company, date_str, e["name"], e["ifsc"], e["account"], e["amount"], e["remarks"])
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=i, column=col_idx, value=val)
            st = BANK_COLUMN_STYLES[col_idx]
            cell.font = st["font"]
            cell.border = st["border"]
            cell.alignment = st["alignment"]
            cell.number_format = st["number_format"]
        ws.row_dimensions[i].height = BANK_ROW_HEIGHT

    for col_letter, width in BANK_COLUMN_WIDTHS.get(company, {}).items():
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def read_bank_employees_from_template(file_storage):
    """Read K (name), M (ifsc), N (account) from a bulk-upload template file.
    Also tries to detect the company from column A of the first data row."""
    wb = load_workbook(file_storage, data_only=True)
    ws = wb.active
    employees = []
    detected_company = None
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row is None or len(row) < 14:
            continue
        code = row[0]
        name = row[10]
        ifsc = row[12]
        account = row[13]
        if not name or not ifsc or not account:
            continue
        name = str(name).strip()
        ifsc = str(ifsc).strip()
        account = str(account).strip()
        employees.append((name, ifsc, account))
        if detected_company is None and code:
            for comp, cfg in BANK_COMPANY_CONFIG.items():
                if str(code).strip() == cfg["code"]:
                    detected_company = comp
    return employees, detected_company


# ============================================================================
# ================  Payroll Processing System  ==============================
# ============================================================================

PAYROLL_MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

PAYROLL_RESULT_COLUMNS = [
    "Name", "Employee Code", "Office", "CTC", "Month Days", "Present Days", "Gross",
    "Basic+DA", "HRA", "Other Allowance", "Total Earning",
    "Employee PF", "Employee ESIC", "PT", "Total Deduction", "Net Pay",
    "Employer PF", "Employer ESIC", "LWF", "Cost To Company", "Additional CTC",
]

# Maps the display column name (used in the results/records UI and Excel export)
# to the payroll_run_records DB column name.
PAYROLL_COLUMN_DB_MAP = {
    "Name": "name", "Employee Code": "code", "Office": "office", "CTC": "ctc",
    "Month Days": "month_days", "Present Days": "present_days", "Gross": "gross",
    "Basic+DA": "basic_da", "HRA": "hra", "Other Allowance": "other_allowance",
    "Total Earning": "total_earning", "Employee PF": "employee_pf",
    "Employee ESIC": "employee_esic", "PT": "pt", "Total Deduction": "total_deduction",
    "Net Pay": "net_pay", "Employer PF": "employer_pf", "Employer ESIC": "employer_esic",
    "LWF": "lwf", "Cost To Company": "cost_to_company", "Additional CTC": "additional_ctc",
}


def calculate_master_payroll(name, ctc, month_days, present_days, code="", office=""):
    """Mirrors the Excel formulas from the original Payroll Processing System exactly."""
    ctc = float(ctc)
    month_days = float(month_days)
    present_days = float(present_days)

    gross_for_pf = ctc / month_days * present_days
    half_gross = gross_for_pf / 2

    if half_gross >= 15000:
        gross = gross_for_pf - 1800
    else:
        gross = gross_for_pf - (half_gross * 0.12)

    basic = gross_for_pf / 2
    hra = gross_for_pf * 0.30
    other = gross - basic - hra

    employee_pf = basic * 0.12 if basic <= 15000 else 1800
    employee_esic = gross * 0.0075 if gross <= 21000 else 0

    total_earning_check = basic + hra + other - employee_pf - employee_esic
    pt = 200 if total_earning_check >= 12000 else 0

    employer_pf = half_gross * 0.12 if half_gross <= 15000 else 1800
    employer_esic = gross * 0.0325 if gross <= 21000 else 0

    total_earning = basic + hra + other
    total_deduction = employee_pf + employee_esic + pt
    net_pay = total_earning - total_deduction

    pf_admin = basic * 0.01 if basic <= 15000 else 15000 * 0.01
    lwf = 24
    additional_ctc = employer_esic + pf_admin + lwf

    total_employer_contribution = employer_pf + employer_esic
    cost_to_company = net_pay + total_deduction + total_employer_contribution

    return {
        "Name": name,
        "Employee Code": code,
        "Office": office,
        "CTC": round(ctc, 2),
        "Month Days": int(month_days),
        "Present Days": int(present_days),
        "Gross": round(gross, 2),
        "Basic+DA": round(basic, 2),
        "HRA": round(hra, 2),
        "Other Allowance": round(other, 2),
        "Total Earning": round(total_earning, 2),
        "Employee PF": round(employee_pf, 2),
        "Employee ESIC": round(employee_esic, 2),
        "PT": round(pt, 2),
        "Total Deduction": round(total_deduction, 2),
        "Net Pay": round(net_pay, 2),
        "Employer PF": round(employer_pf, 2),
        "Employer ESIC": round(employer_esic, 2),
        "LWF": round(lwf, 2),
        "Cost To Company": round(cost_to_company, 2),
        "Additional CTC": round(additional_ctc, 2),
    }


def write_payroll_excel(columns, records):
    """records: list of dicts keyed by display column names (PAYROLL_RESULT_COLUMNS,
    optionally prefixed with Month/Year). Returns a BytesIO xlsx file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"

    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, r in enumerate(records, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=r.get(col_name, ""))

    for col_idx, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(col_name) + 2)

    ws.freeze_panes = "A2"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def payroll_record_row_to_dict(row):
    """Converts a payroll_run_records DB row (+ month/year/batch_no) into the display-column dict."""
    d = dict(row)
    out = {"id": d.get("id"), "Month": d.get("month", ""), "Year": d.get("year", ""), "Batch No": d.get("batch_no", "")}
    for display_col, db_col in PAYROLL_COLUMN_DB_MAP.items():
        out[display_col] = d.get(db_col)
    return out


PAYROLL_MONTH_ABBREV = {
    "January": "Jan", "February": "Feb", "March": "Mar", "April": "Apr",
    "May": "May", "June": "Jun", "July": "Jul", "August": "Aug",
    "September": "Sep", "October": "Oct", "November": "Nov", "December": "Dec",
}


def next_payroll_batch_no(cursor, month, year):
    """Batch No format: MMM-YY/NN — serial resets to 01 for every new month/year."""
    cursor.execute("SELECT COUNT(*) AS c FROM payroll_runs WHERE month = ? AND year = ?", (month, year))
    count = cursor.fetchone()["c"]
    abbr = PAYROLL_MONTH_ABBREV.get(month, str(month)[:3])
    yy = str(year)[-2:]
    return f"{abbr}-{yy}/{str(count + 1).zfill(2)}"


def apply_weekend_absence_rule(records, start_date=None, end_date=None):
    if not records:
        return records

    def normalize_date(value):
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except Exception:
            return None

    def weekday_score(status):
        if status in ('P', 'H', 'WO'):
            return 1.0
        if status in ('AH', 'HF'):
            return 0.5
        return 0.0

    if start_date:
        try:
            start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
        except Exception:
            start_dt = None
    else:
        start_dt = None
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, '%Y-%m-%d').date()
        except Exception:
            end_dt = None
    else:
        end_dt = None

    by_employee = {}
    for row in records:
        emp_id = row.get('employee_id')
        if emp_id is None:
            continue
        dt = normalize_date(row.get('date'))
        if not dt:
            continue
        week_key = (dt.isocalendar()[0], dt.isocalendar()[1])
        by_employee.setdefault(emp_id, {})
        by_employee[emp_id].setdefault(week_key, {'records': {}})
        by_employee[emp_id][week_key]['records'][dt] = row

    extra_rows = []
    for emp_id, weeks in by_employee.items():
        for week_key, week_data in weeks.items():
            year, week_num = week_key
            try:
                monday = datetime.fromisocalendar(year, week_num, 1).date()
            except ValueError:
                continue
            total_score = 0.0
            for offset in range(5):
                day = monday + timedelta(days=offset)
                row = week_data['records'].get(day)
                status = row.get('status') if row else None
                total_score += weekday_score(status)
            if total_score >= 3.5:
                continue
            for offset in (5, 6):
                day = monday + timedelta(days=offset)
                if start_dt and day < start_dt:
                    continue
                if end_dt and day > end_dt:
                    continue
                existing = week_data['records'].get(day)
                if existing:
                    continue
                extra_rows.append({
                    'id': None,
                    'employee_id': emp_id,
                    'date': day.strftime('%Y-%m-%d'),
                    'status': 'A',
                    'shift_name': None,
                    'check_in': None,
                    'check_out': None
                })

    # Additional conditions for WO and H to A
    for emp_id, weeks in by_employee.items():
        week_keys = sorted(weeks.keys())
        for i, week_key in enumerate(week_keys):
            year, week_num = week_key
            monday = datetime.fromisocalendar(year, week_num, 1).date()
            friday = monday + timedelta(days=4)
            sat = monday + timedelta(days=5)
            sun = monday + timedelta(days=6)
            friday_row = weeks[week_key]['records'].get(friday)
            friday_status = friday_row.get('status') if friday_row else None
            next_monday = monday + timedelta(days=7)
            next_week_key = (next_monday.isocalendar()[0], next_monday.isocalendar()[1])
            next_monday_row = weeks.get(next_week_key, {}).get('records', {}).get(next_monday)
            next_monday_status = next_monday_row.get('status') if next_monday_row else None

            # Condition 1: absent on Friday and Monday -> WO to A
            if friday_status == 'A' and next_monday_status == 'A':
                pass  # condition for auto, but do not override existing records; respect manual marks from pop-up

            # Condition 2: holiday on Friday with absences around -> WO and H to A
            if friday_status == 'H':
                thursday = monday + timedelta(days=3)
                thursday_row = weeks[week_key]['records'].get(thursday)
                thursday_status = thursday_row.get('status') if thursday_row else None
                tuesday = next_monday + timedelta(days=1)
                tuesday_row = weeks.get(next_week_key, {}).get('records', {}).get(tuesday)
                tuesday_status = tuesday_row.get('status') if tuesday_row else None
                absent_around = (thursday_status == 'A' or friday_status == 'A') and (next_monday_status == 'A' or tuesday_status == 'A')
                if absent_around:
                    pass  # condition met, but respect any explicit status on sat/sun (manual override allowed)

            # Condition 3: holiday on Monday with absences around -> WO and H to A
            if next_monday_status == 'H':
                saturday = monday + timedelta(days=5)
                saturday_row = weeks[week_key]['records'].get(saturday)
                saturday_status = saturday_row.get('status') if saturday_row else None
                tuesday = next_monday + timedelta(days=1)
                tuesday_row = weeks.get(next_week_key, {}).get('records', {}).get(tuesday)
                tuesday_status = tuesday_row.get('status') if tuesday_row else None
                wednesday = next_monday + timedelta(days=2)
                wednesday_row = weeks.get(next_week_key, {}).get('records', {}).get(wednesday)
                wednesday_status = wednesday_row.get('status') if wednesday_row else None
                absent_around = (friday_status == 'A' or saturday_status == 'A') and (tuesday_status == 'A' or wednesday_status == 'A')
                if absent_around:
                    pass  # condition met, but respect any explicit status on sat/sun (manual override allowed)

    return records + extra_rows


def is_late_arrival(check_in):
    """Late coming if check_in is between 19:40 and 21:30 (for the 3 allowed lates policy)."""
    if not check_in:
        return False
    try:
        parts = str(check_in).split(':')
        h = int(parts[0])
        m = int(parts[1]) if len(parts) > 1 else 0
        cin_min = h * 60 + m
        late_start = 19 * 60 + 40
        late_end = 21 * 60 + 30
        if late_start <= cin_min <= late_end:
            return True
    except Exception:
        pass
    return False


def is_half_day_by_time(check_in, check_out):
    """Time Policy:
    - in on time (before 19:40) and out 22:00-23:59 → AH
    - in 22:00-01:00 and out on time (before ~20:00) → AH
    """
    if not check_in or not check_out:
        return False
    try:
        cin_h, cin_m = [int(x) for x in str(check_in).split(':')[:2]]
        cout_h, cout_m = [int(x) for x in str(check_out).split(':')[:2]]
        cin_minutes = cin_h * 60 + cin_m
        cout_minutes = cout_h * 60 + cout_m

        # 1. in on time AND out late evening
        in_on_time = cin_minutes < (19 * 60 + 40)
        out_late_window = (22 * 60) <= cout_minutes <= (23 * 60 + 59)
        if in_on_time and out_late_window:
            return True

        # 2. in late night 22:00-01:00 AND out on time
        in_late_night = (22 * 60 <= cin_minutes) or (cin_minutes <= 60)  # up to 1:00
        out_on_time = cout_minutes < (20 * 60)
        if in_late_night and out_on_time:
            return True
    except Exception:
        pass
    return False


def get_shift_end_time(shift_id):
    if not shift_id:
        return None
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT end_time FROM shifts WHERE id = ?", (shift_id,))
    row = cursor.fetchone()
    conn.close()
    if row:
        if isinstance(row, (list, tuple)):
            return row[0]
        else:
            return row['end_time'] if 'end_time' in row.keys() else row[0]
    return None


def is_out_on_schedule(check_out, shift_id=None):
    """Returns True if check_out is on or before the shift's scheduled end time (or True if no shift)."""
    if not check_out:
        return True  # if no out time, assume on schedule for late policy
    if not shift_id:
        return True
    shift_end = get_shift_end_time(shift_id)
    if not shift_end:
        return True
    try:
        def to_minutes(t):
            h, m = [int(x) for x in str(t).split(':')[:2]]
            return h * 60 + m
        out_min = to_minutes(check_out)
        end_min = to_minutes(shift_end)
        return out_min <= end_min
    except Exception:
        return True


def count_late_arrivals_in_month(employee_id, year, month, before_date=None):
    """Count how many qualified late comings (in 19:40-21:30 window AND out on schedule) for emp in given year-month."""
    conn = get_connection()
    cursor = conn.cursor()
    q = """SELECT check_in, check_out, shift_id FROM attendance
           WHERE employee_id = ?
             AND strftime('%Y', date) = ?
             AND strftime('%m', date) = ?"""
    p = [employee_id, str(year), f"{month:02d}"]
    if before_date:
        q += " AND date < ?"
        p.append(before_date)
    cursor.execute(q, p)
    rows = dict_rows(cursor.fetchall())
    conn.close()
    cnt = 0
    for r in rows:
        cin = r.get('check_in')
        cout = r.get('check_out')
        sid = r.get('shift_id')
        if is_late_arrival(cin) and is_out_on_schedule(cout, sid):
            cnt += 1
    return cnt


def get_effective_status_for_save(employee_id, date_str, submitted_status, check_in, check_out, shift_id=None):
    """Apply late coming and time policies to decide the status that should be saved."""
    effective = submitted_status or 'P'
    # Time policy (takes precedence if matches)
    if check_in and check_out and is_half_day_by_time(check_in, check_out):
        return 'AH'
    # Late coming: in window 19:40-21:30 AND out on schedule
    if check_in and is_late_arrival(check_in) and is_out_on_schedule(check_out, shift_id):
        try:
            dt = datetime.strptime(date_str, '%Y-%m-%d')
            prev = count_late_arrivals_in_month(employee_id, dt.year, dt.month, before_date=date_str)
            if prev >= 3:
                effective = 'AH'
        except Exception:
            pass
    return effective


def get_probation_end_date(joining_str):
    if not joining_str:
        return None
    try:
        jdate = datetime.strptime(str(joining_str).strip(), '%Y-%m-%d').date()
        new_month = jdate.month + 3
        new_year = jdate.year
        if new_month > 12:
            new_year += 1
            new_month -= 12
        try:
            return date(new_year, new_month, jdate.day)
        except ValueError:
            # last day of the month
            if new_month == 12:
                next_month = date(new_year + 1, 1, 1)
            else:
                next_month = date(new_year, new_month + 1, 1)
            return next_month - timedelta(days=1)
    except Exception:
        return None


def is_eligible_for_earned_leave(employee_id, year, month, cursor):
    """Returns True if employee has completed 3-month probation by end of the given month."""
    cursor.execute("SELECT joining_date FROM employees WHERE id = ?", (employee_id,))
    row = cursor.fetchone()
    if not row:
        return False
    if isinstance(row, (list, tuple)):
        joining = row[0]
    else:
        joining = row["joining_date"] if "joining_date" in row.keys() else row[0]
    probation_end = get_probation_end_date(joining)
    if not probation_end:
        return False
    if month == 12:
        month_end = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        month_end = date(year, month + 1, 1) - timedelta(days=1)
    return probation_end <= month_end


def ensure_leave_balance(cursor, employee_id, year):
    cursor.execute(
        "INSERT OR IGNORE INTO leave_balances (employee_id, earned_leave, special_leave, loss_of_pay, year) VALUES (?, 0, 0, 0, ?)",
        (employee_id, year)
    )


def adjust_leave_balance(cursor, employee_id, leave_type, start_date_str, end_date_str, reverse=False):
    """Deduct (or reverse) days from the appropriate balance bucket."""
    try:
        start_d = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_d = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        days = (end_d - start_d).days + 1
        y = start_d.year
        ensure_leave_balance(cursor, employee_id, y)
        if leave_type == "Earned Leave":
            delta = days if reverse else -days
            cursor.execute(
                "UPDATE leave_balances SET earned_leave = earned_leave + ? WHERE employee_id = ? AND year = ?",
                (delta, employee_id, y)
            )
        elif leave_type == "Special Leave":
            delta = days if reverse else -days
            cursor.execute(
                "UPDATE leave_balances SET special_leave = special_leave + ? WHERE employee_id = ? AND year = ?",
                (delta, employee_id, y)
            )
        elif leave_type == "Loss of Pay":
            delta = -days if reverse else days
            cursor.execute(
                "UPDATE leave_balances SET loss_of_pay = loss_of_pay + ? WHERE employee_id = ? AND year = ?",
                (delta, employee_id, y)
            )
    except Exception:
        pass


def mark_attendance_for_leave_days(cursor, employee_id, start_date_str, end_date_str, leave_type, reason):
    """Mark each day in the leave range in attendance as 'A' (so it affects earned accrual present count) with remarks."""
    try:
        start = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        current = start
        status_for_att = 'A'
        rem = f"{leave_type} - {reason or ''}".strip()[:200]
        while current <= end:
            dstr = current.strftime('%Y-%m-%d')
            cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (employee_id, dstr))
            existing = cursor.fetchone()
            if existing:
                cursor.execute("UPDATE attendance SET status = ?, remarks = ? WHERE id = ?", (status_for_att, rem, existing[0]))
            else:
                cursor.execute(
                    "INSERT INTO attendance (employee_id, date, status, remarks) VALUES (?, ?, ?, ?)",
                    (employee_id, dstr, status_for_att, rem)
                )
            current += timedelta(days=1)
    except Exception:
        pass


def calculate_pay_components(ctc, pf_active=True, esic_active=True, professional_tax_active=True, lwf_active=True):
    basic = round(ctc * 0.30, 2)
    da = round(ctc * 0.20, 2)
    hra = round(ctc * 0.30, 2)
    conveyance = round(ctc * 0.10, 2)
    special_allowance = round(ctc * 0.10, 2)
    gross = round(basic + da + hra + conveyance + special_allowance, 2)
    pf = round(basic * 0.12, 2) if pf_active else 0.0
    esic = round(gross * 0.0475, 2) if esic_active else 0.0
    professional_tax = 200.0 if professional_tax_active else 0.0
    lwf = 30.0 if lwf_active else 0.0
    total_deductions = round(pf + esic + professional_tax + lwf, 2)
    net_pay = round(gross - total_deductions, 2)
    return {
        "basic": basic,
        "da": da,
        "hra": hra,
        "conveyance": conveyance,
        "special_allowance": special_allowance,
        "gross": gross,
        "pf": pf,
        "esic": esic,
        "professional_tax": professional_tax,
        "lwf": lwf,
        "total_deductions": total_deductions,
        "net_pay": net_pay
    }


def get_employee_list():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT e.*, d.name AS department_name, des.title AS designation_title, o.name AS office_name, m.first_name || ' ' || m.last_name AS manager_name "
        "FROM employees e "
        "LEFT JOIN departments d ON e.department_id = d.id "
        "LEFT JOIN designations des ON e.designation_id = des.id "
        "LEFT JOIN offices o ON e.office_id = o.id "
        "LEFT JOIN employees m ON e.manager_id = m.id "
        "ORDER BY e.first_name, e.last_name"
    )
    rows = cursor.fetchall()
    conn.close()
    return dict_rows(rows)


def get_lookup_data():
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM departments ORDER BY name")
    departments = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM offices ORDER BY name")
    offices = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM designations ORDER BY title")
    designations = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM shifts ORDER BY name")
    shifts = dict_rows(cursor.fetchall())
    employees = get_employee_list()

    conn.close()
    return {
        "departments": departments,
        "offices": offices,
        "designations": designations,
        "shifts": shifts,
        "employees": employees,
        "leave_types": LEAVE_TYPES
    }

@app.route("/")
def index():
    if not session.get('logged_in'):
        return redirect(url_for('login_page'))
    return render_template("index.html")

@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "POST":
        user_id = request.form.get("user_id", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if user_id == ADMIN_USER and password == ADMIN_PASSWORD:
            session['logged_in'] = True
            session['is_admin'] = True
            session['employee_id'] = None
            return jsonify({"success": True})

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM employees WHERE phone = ?", (user_id,))
        emp = cursor.fetchone()
        conn.close()

        if not emp:
            return jsonify({"success": False, "error": "Employee not found with this mobile number."})

        stored_pass = emp['password'] or ""
        if not stored_pass:
            # First time login - set the password (hashed)
            if confirm and password != confirm:
                return jsonify({"success": False, "error": "Passwords do not match.", "first_time": True})
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("UPDATE employees SET password = ? WHERE id = ?", (generate_password_hash(password), emp['id']))
            conn.commit()
            conn.close()
            session['logged_in'] = True
            session['is_admin'] = False
            session['employee_id'] = emp['id']
            session['user_phone'] = user_id
            return jsonify({"success": True})

        # Support both new hashed passwords and any pre-existing plaintext ones
        # (old rows get upgraded to a hash automatically on next successful login)
        password_ok = False
        try:
            password_ok = check_password_hash(stored_pass, password)
        except ValueError:
            password_ok = (stored_pass == password)
            if password_ok:
                conn = get_connection()
                cursor = conn.cursor()
                cursor.execute("UPDATE employees SET password = ? WHERE id = ?", (generate_password_hash(password), emp['id']))
                conn.commit()
                conn.close()

        if password_ok:
            session['logged_in'] = True
            session['is_admin'] = False
            session['employee_id'] = emp['id']
            session['user_phone'] = user_id
            return jsonify({"success": True})
        else:
            return jsonify({"success": False, "error": "Invalid password."})

    # GET - show login page
    if session.get('logged_in'):
        return redirect(url_for('index'))
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('login_page'))

@app.route("/api/current_user")
def current_user():
    if not session.get('logged_in'):
        return jsonify({"error": "Not logged in"}), 401
    if session.get('is_admin'):
        return jsonify({"is_admin": True, "employee_id": None})
    emp_id = session.get('employee_id')
    if emp_id:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, employee_code, first_name, last_name, phone FROM employees WHERE id = ?", (emp_id,))
        emp = cursor.fetchone()
        conn.close()
        if emp:
            return jsonify({
                "is_admin": False,
                "employee_id": emp['id'],
                "phone": emp['phone'],
                "name": f"{emp['first_name']} {emp['last_name']}",
                "employee_code": emp['employee_code']
            })
    return jsonify({"is_admin": False, "employee_id": None})

@app.route("/api/lookups")
def lookups():
    return jsonify(get_lookup_data())

@app.route("/api/departments", methods=["GET", "POST"])
def manage_departments():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM departments ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    description = payload.get("description", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Department name is required."}), 400
    try:
        cursor.execute("INSERT INTO departments (name, description) VALUES (?, ?)", (name, description))
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Department already exists."}), 400
    finally:
        conn.close()

@app.route("/api/departments/<int:department_id>", methods=["DELETE"])
def delete_department(department_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM departments WHERE id = ?", (department_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/offices", methods=["GET", "POST"])
def manage_offices():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM offices ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    location = payload.get("location", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Office name is required."}), 400
    try:
        cursor.execute("INSERT INTO offices (name, location) VALUES (?, ?)", (name, location))
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Office already exists."}), 400
    finally:
        conn.close()

@app.route("/api/offices/<int:office_id>", methods=["DELETE"])
def delete_office(office_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM offices WHERE id = ?", (office_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/designations", methods=["GET", "POST"])
def manage_designations():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM designations ORDER BY title")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    title = payload.get("title", "").strip()
    department_id = payload.get("department_id")
    if not title:
        conn.close()
        return jsonify({"error": "Designation title is required."}), 400
    cursor.execute("INSERT INTO designations (title, department_id) VALUES (?, ?)", (title, department_id))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/designations/<int:designation_id>", methods=["DELETE"])
def delete_designation(designation_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM designations WHERE id = ?", (designation_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/shifts", methods=["GET", "POST"])
def manage_shifts():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM shifts ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    start_time = payload.get("start_time", "").strip()
    end_time = payload.get("end_time", "").strip()
    description = payload.get("description", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Shift name is required."}), 400
    cursor.execute(
        "INSERT INTO shifts (name, start_time, end_time, description) VALUES (?, ?, ?, ?)",
        (name, start_time, end_time, description)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/shifts/<int:shift_id>", methods=["GET", "PUT", "DELETE"])
def manage_shift(shift_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM shifts WHERE id = ?", (shift_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    if request.method == "PUT":
        payload = request.json or {}
        name = payload.get("name", "").strip()
        start_time = payload.get("start_time", "").strip()
        end_time = payload.get("end_time", "").strip()
        description = payload.get("description", "").strip()
        if not name:
            conn.close()
            return jsonify({"error": "Shift name is required."}), 400
        cursor.execute(
            "UPDATE shifts SET name = ?, start_time = ?, end_time = ?, description = ? WHERE id = ?",
            (name, start_time, end_time, description, shift_id)
        )
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    cursor.execute("DELETE FROM shifts WHERE id = ?", (shift_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/employees", methods=["GET", "POST"])
def manage_employees():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        data = get_employee_list()
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    required_fields = ["employee_code", "first_name", "last_name", "department_id", "designation_id", "office_id", "joining_date", "ctc"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    employee_data = (
        payload.get("employee_code", "").strip() or None,
        payload.get("first_name").strip(),
        payload.get("last_name").strip(),
        payload.get("email", "").strip(),
        payload.get("phone", "").strip(),
        payload.get("dob", "").strip() or None,
        payload.get("gender", "").strip(),
        payload.get("department_id"),
        payload.get("designation_id"),
        payload.get("office_id"),
        payload.get("manager_id"),
        payload.get("joining_date", "").strip() or None,
        float(payload.get("ctc") or 0),
        payload.get("status", "Active").strip(),
        payload.get("address", "").strip(),
        int(bool(payload.get("pf_active", 1))),
        int(bool(payload.get("esic_active", 1))),
        int(bool(payload.get("professional_tax_active", 1))),
        int(bool(payload.get("lwf_active", 1)))
    )

    try:
        cursor.execute(
            "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            employee_data
        )
        employee_id = cursor.lastrowid
        # Insert default leave balance
        current_year = datetime.now().year
        cursor.execute(
            "INSERT INTO leave_balances (employee_id, earned_leave, special_leave, loss_of_pay, year) VALUES (?, 0, 0, 0, ?)",
            (employee_id, current_year)
        )
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Employee code must be unique."}), 400
    finally:
        conn.close()

@app.route("/api/employees/<int:employee_id>", methods=["GET", "PUT", "DELETE"])
def employee_detail(employee_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM employees WHERE id = ?", (employee_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    if request.method == "DELETE":
        cursor.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    payload = request.json or {}
    employee_data = (
        payload.get("employee_code", "").strip(),
        payload.get("first_name", "").strip(),
        payload.get("last_name", "").strip(),
        payload.get("email", "").strip(),
        payload.get("phone", "").strip(),
        payload.get("dob", "").strip(),
        payload.get("gender", "").strip(),
        payload.get("department_id"),
        payload.get("designation_id"),
        payload.get("office_id"),
        payload.get("manager_id"),
        payload.get("joining_date", "").strip(),
        float(payload.get("ctc", 0)),
        payload.get("status", "Active").strip(),
        payload.get("address", "").strip(),
        int(bool(payload.get("pf_active", 1))),
        int(bool(payload.get("esic_active", 1))),
        int(bool(payload.get("professional_tax_active", 1))),
        int(bool(payload.get("lwf_active", 1))),
        employee_id
    )
    cursor.execute(
        "UPDATE employees SET employee_code = ?, first_name = ?, last_name = ?, email = ?, phone = ?, dob = ?, gender = ?, department_id = ?, designation_id = ?, office_id = ?, manager_id = ?, joining_date = ?, ctc = ?, status = ?, address = ?, pf_active = ?, esic_active = ?, professional_tax_active = ?, lwf_active = ? WHERE id = ?",
        employee_data
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/employees/<int:employee_id>/statutory", methods=["PUT"])
def update_statutory(employee_id):
    conn = get_connection()
    cursor = conn.cursor()
    payload = request.json or {}
    updates = {}
    if "pf_active" in payload:
        updates["pf_active"] = int(bool(payload["pf_active"]))
    if "esic_active" in payload:
        updates["esic_active"] = int(bool(payload["esic_active"]))
    if "pt_active" in payload:
        updates["pt_active"] = int(bool(payload["pt_active"]))
    if "lwf_active" in payload:
        updates["lwf_active"] = int(bool(payload["lwf_active"]))
    if not updates:
        conn.close()
        return jsonify({"error": "No valid fields to update."}), 400
    set_clause = ", ".join(f"{k} = ?" for k in updates)
    values = list(updates.values()) + [employee_id]
    cursor.execute(f"UPDATE employees SET {set_clause} WHERE id = ?", values)
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/download-template")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = ["employee_code", "first_name", "last_name", "email", "phone", "dob", "gender", "department_name", "designation_title", "office_name", "manager_name", "joining_date", "ctc", "status", "address"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    sample_data = ["EMP001", "John", "Doe", "john@example.com", "1234567890", "1990-01-01", "Male", "Engineering", "Developer", "Head Office", "", "2023-01-01", "50000", "Active", "123 Main St"]
    for col_num, value in enumerate(sample_data, 1):
        ws.cell(row=2, column=col_num, value=value)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="employee_import_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/api/employees/bulk-import", methods=["POST"])
def bulk_import_employees():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Invalid file type. Please upload an Excel file."}), 400
    
    try:
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File is empty"}), 400
        headers = [str(h).strip().lower() for h in rows[0]]
        expected_headers = ["employee_code", "first_name", "last_name", "email", "phone", "dob", "gender", "department_name", "designation_title", "office_name", "manager_name", "joining_date", "ctc", "status", "address"]
        if headers != expected_headers:
            return jsonify({"error": "Invalid headers. Please use the template."}), 400
        
        conn = get_connection()
        cursor = conn.cursor()
        imported = 0
        updated = 0
        for row in rows[1:]:
            # Safe access to avoid index errors on short rows
            def g(i, default=None):
                try:
                    return row[i]
                except (IndexError, TypeError):
                    return default

            def parse_str(val):
                s = str(val or '').strip()
                return s if s else None

            emp_code_raw = g(0)
            emp_code = parse_str(emp_code_raw)
            if not emp_code:
                continue

            first = parse_str(g(1))
            last = parse_str(g(2))
            email = parse_str(g(3))
            phone = parse_str(g(4))
            dob = normalize_date_value(g(5))
            gender = parse_str(g(6))
            dept_name = parse_str(g(7))
            des_title = parse_str(g(8))
            off_name = parse_str(g(9))
            mgr_raw = parse_str(g(10))
            joining_date = normalize_date_value(g(11))
            ctc_raw = g(12)
            ctc_val = None
            if ctc_raw is not None:
                try:
                    ctc_str = str(ctc_raw).replace(',', '').replace('$', '').strip()
                    ctc_val = float(ctc_str) if ctc_str else None
                except Exception:
                    ctc_val = None
            status = parse_str(g(13))
            address = parse_str(g(14))

            # Check if employee exists by code
            cursor.execute("SELECT id FROM employees WHERE employee_code = ?", (emp_code,))
            existing = cursor.fetchone()

            if existing:
                # UPDATE only non-blank fields from excel; leave DB values if blank in excel
                set_clauses = []
                params = []

                if first is not None:
                    set_clauses.append("first_name=?")
                    params.append(first)
                if last is not None:
                    set_clauses.append("last_name=?")
                    params.append(last)
                if email is not None:
                    set_clauses.append("email=?")
                    params.append(email)
                if phone is not None:
                    set_clauses.append("phone=?")
                    params.append(phone)
                if dob is not None:
                    set_clauses.append("dob=?")
                    params.append(dob)
                if gender is not None:
                    set_clauses.append("gender=?")
                    params.append(gender)

                if dept_name is not None:
                    department_id = None
                    if dept_name:
                        cursor.execute("SELECT id FROM departments WHERE name = ?", (dept_name,))
                        dep = cursor.fetchone()
                        department_id = dep[0] if dep else None
                    set_clauses.append("department_id=?")
                    params.append(department_id)

                if des_title is not None:
                    designation_id = None
                    if des_title:
                        cursor.execute("SELECT id FROM designations WHERE title = ?", (des_title,))
                        des = cursor.fetchone()
                        designation_id = des[0] if des else None
                    set_clauses.append("designation_id=?")
                    params.append(designation_id)

                if off_name is not None:
                    office_id = None
                    if off_name:
                        cursor.execute("SELECT id FROM offices WHERE name = ?", (off_name,))
                        off = cursor.fetchone()
                        office_id = off[0] if off else None
                    set_clauses.append("office_id=?")
                    params.append(office_id)

                if mgr_raw is not None:
                    manager_id = None
                    if mgr_raw:
                        names = mgr_raw.split()
                        if len(names) >= 2:
                            cursor.execute("SELECT id FROM employees WHERE first_name = ? AND last_name = ?", (names[0], names[1]))
                            mgr = cursor.fetchone()
                            manager_id = mgr[0] if mgr else None
                    set_clauses.append("manager_id=?")
                    params.append(manager_id)

                if joining_date is not None:
                    set_clauses.append("joining_date=?")
                    params.append(joining_date)
                if ctc_val is not None:
                    set_clauses.append("ctc=?")
                    params.append(ctc_val)
                if status is not None:
                    set_clauses.append("status=?")
                    params.append(status)
                if address is not None:
                    set_clauses.append("address=?")
                    params.append(address)

                if set_clauses:
                    params.append(emp_code)
                    cursor.execute(
                        f"UPDATE employees SET {', '.join(set_clauses)} WHERE employee_code = ?",
                        params
                    )
                    updated += 1
            else:
                # INSERT new: require first + last (as before)
                if first is None or last is None:
                    continue

                department_id = None
                if dept_name:
                    cursor.execute("SELECT id FROM departments WHERE name = ?", (dept_name,))
                    dep = cursor.fetchone()
                    department_id = dep[0] if dep else None

                designation_id = None
                if des_title:
                    cursor.execute("SELECT id FROM designations WHERE title = ?", (des_title,))
                    des = cursor.fetchone()
                    designation_id = des[0] if des else None

                office_id = None
                if off_name:
                    cursor.execute("SELECT id FROM offices WHERE name = ?", (off_name,))
                    off = cursor.fetchone()
                    office_id = off[0] if off else None

                manager_id = None
                if mgr_raw:
                    names = mgr_raw.split()
                    if len(names) >= 2:
                        cursor.execute("SELECT id FROM employees WHERE first_name = ? AND last_name = ?", (names[0], names[1]))
                        mgr = cursor.fetchone()
                        manager_id = mgr[0] if mgr else None

                gender = gender or "Male"
                status = status or "Active"
                ctc_val = ctc_val if ctc_val is not None else 0.0

                employee_data = (
                    emp_code,
                    first,
                    last,
                    email,
                    phone,
                    dob,
                    gender,
                    department_id,
                    designation_id,
                    office_id,
                    manager_id,
                    joining_date,
                    ctc_val,
                    status,
                    address,
                    1, 1, 1, 1
                )
                try:
                    cursor.execute(
                        "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        employee_data
                    )
                    employee_id = cursor.lastrowid
                    current_year = datetime.now().year
                    cursor.execute(
                        "INSERT OR IGNORE INTO leave_balances (employee_id, earned_leave, special_leave, loss_of_pay, year) VALUES (?, 0, 0, 0, ?)",
                        (employee_id, current_year)
                    )
                    imported += 1
                except sqlite3.IntegrityError:
                    continue
        conn.commit()
        conn.close()
        return jsonify({"imported": imported, "updated": updated})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/employee-joining-forms", methods=["POST"])
def save_employee_joining_form():
    conn = get_connection()
    cursor = conn.cursor()

    # Generate employee code first
    employee_code = get_next_employee_code()

    # Text fields from form
    full_name = request.form.get("joining_full_name", "")
    contact_no = request.form.get("joining_contact_no", "")
    email_id = request.form.get("joining_email_id", "")
    designation = request.form.get("joining_designation", "")
    joining_date = request.form.get("joining_joining_date", "")
    permanent_address = request.form.get("joining_permanent_address", "")
    date_of_birth = request.form.get("joining_date_of_birth", "")
    gender = request.form.get("joining_gender", "")
    marital_status = request.form.get("joining_marital_status", "")
    pan_no = request.form.get("joining_pan_no", "")
    aadhar_no = request.form.get("joining_aadhar_no", "")
    bank_name = request.form.get("joining_bank_name", "")
    bank_account_number = request.form.get("joining_bank_account_number", "")
    ifsc_code = request.form.get("joining_ifsc_code", "")
    branch_name = request.form.get("joining_branch_name", "")
    highest_qualification = request.form.get("joining_highest_qualification", "")
    university_board = request.form.get("joining_university_board", "")
    year_of_passing = request.form.get("joining_year_of_passing", "")
    percentage_grade = request.form.get("joining_percentage_grade", "")
    emergency_person_name = request.form.get("joining_emergency_person_name", "")
    emergency_person_mobile = request.form.get("joining_emergency_person_mobile", "")
    emergency_relation = request.form.get("joining_emergency_relation", "")
    emergency_person_address = request.form.get("joining_emergency_person_address", "")
    prev_company_name = request.form.get("joining_prev_company_name", "")
    prev_designation = request.form.get("joining_prev_designation", "")
    prev_duration = request.form.get("joining_prev_duration", "")
    prev_last_salary = request.form.get("joining_prev_last_salary", "")

    # Handle file uploads (local + optional Google Drive)
    def save_uploaded_file(file_key):
        if file_key in request.files:
            file = request.files[file_key]
            if file and file.filename:
                filename = secure_filename(file.filename)
                filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
                filepath = os.path.join(UPLOAD_FOLDER, filename)
                file.save(filepath)
                
                # Optional Google Drive upload (uses your existing google_drive_manager setup)
                try:
                    from google_drive_manager import GoogleDriveManager
                    gdm = GoogleDriveManager()
                    # Using the existing "Offer Letters" folder for now; documents will be private to your account
                    folder_id = gdm.create_offer_letters_folder()
                    with open(filepath, 'rb') as f:
                        content = f.read()
                    drive_res = gdm.upload_file(filename, content, folder_id=folder_id)
                    print(f"✅ Also uploaded to Google Drive: {drive_res.get('view_link')}")
                except Exception as drive_err:
                    print(f"⚠️ Google Drive upload skipped for {file_key}: {drive_err}")
                
                return filename
        return ""
    photo_filename = save_uploaded_file("joining_photo")
    pan_card_filename = save_uploaded_file("joining_pan_card")
    aadhar_card_filename = save_uploaded_file("joining_aadhar_card")
    cheque_passbook_filename = save_uploaded_file("joining_cheque_passbook")
    highest_education_cert_filename = save_uploaded_file("joining_highest_education_cert")
    last_3_month_salary_slip_filename = save_uploaded_file("joining_last_3_month_salary_slip")
    prev_employment_docs_filename = save_uploaded_file("joining_prev_employment_docs")

    cursor.execute(
        "INSERT INTO employee_joining_forms ("
        "full_name, contact_no, email_id, designation, joining_date, permanent_address, date_of_birth, gender, marital_status, "
        "pan_no, aadhar_no, bank_name, bank_account_number, ifsc_code, branch_name, "
        "highest_qualification, university_board, year_of_passing, percentage_grade, "
        "emergency_person_name, emergency_person_mobile, emergency_relation, emergency_person_address, "
        "prev_company_name, prev_designation, prev_duration, prev_last_salary, "
        "photo_filename, pan_card_filename, aadhar_card_filename, cheque_passbook_filename, "
        "highest_education_cert_filename, last_3_month_salary_slip_filename, prev_employment_docs_filename, "
        "employee_code, status, submitted_on"
        ") VALUES ("
        "?, ?, ?, ?, ?, ?, ?, ?, ?, "
        "?, ?, ?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, "
        "?, ?, ?"
        ")",
        (
            full_name, contact_no, email_id, designation, joining_date, permanent_address, date_of_birth, gender, marital_status,
            pan_no, aadhar_no, bank_name, bank_account_number, ifsc_code, branch_name,
            highest_qualification, university_board, year_of_passing, percentage_grade,
            emergency_person_name, emergency_person_mobile, emergency_relation, emergency_person_address,
            prev_company_name, prev_designation, prev_duration, prev_last_salary,
            photo_filename, pan_card_filename, aadhar_card_filename, cheque_passbook_filename,
            highest_education_cert_filename, last_3_month_salary_slip_filename, prev_employment_docs_filename,
            employee_code, 'pending',
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        )
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True, "employee_code": employee_code})


@app.route("/api/reject-joining", methods=["POST"])
def reject_joining():
    payload = request.json or {}
    joining_id = payload.get("id")
    if not joining_id:
        return jsonify({"error": "Joining ID required"}), 400
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE employee_joining_forms SET status = 'rejected' WHERE id = ?",
        (joining_id,)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route("/api/pending-joinings")
def get_pending_joinings():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT * FROM employee_joining_forms 
        WHERE status = 'pending' 
        ORDER BY submitted_on DESC
    """)
    rows = cursor.fetchall()
    conn.close()
    return jsonify([dict(row) for row in rows])


@app.route("/api/approve-joining", methods=["POST"])
def approve_joining():
    payload = request.json or {}
    joining_id = payload.get("id")
    if not joining_id:
        return jsonify({"error": "Joining ID required"}), 400
    
    conn = get_connection()
    cursor = conn.cursor()
    
    # Get the pending record
    cursor.execute("SELECT * FROM employee_joining_forms WHERE id = ?", (joining_id,))
    joining = cursor.fetchone()
    if not joining or joining['status'] != 'pending':
        conn.close()
        return jsonify({"error": "Invalid or already processed joining"}), 400
    
    # Split full name for employees table
    full_name = joining['full_name'] or ''
    parts = full_name.split(' ', 1)
    first_name = parts[0] if parts else ''
    last_name = parts[1] if len(parts) > 1 else ''
    
    # Map fields (some will be NULL/default since joining form has limited data)
    employee_data = (
        joining['employee_code'],
        first_name,
        last_name,
        joining['email_id'],
        joining['contact_no'],
        joining['date_of_birth'],
        joining['gender'],
        None,  # department_id - admin can update later
        None,  # designation_id - use text for now or map
        None,  # office_id
        None,  # manager_id
        joining['joining_date'],
        0,     # ctc default
        'Active',
        joining['permanent_address'],
        1, 1, 1, 1
    )
    
    try:
        cursor.execute(
            "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            employee_data
        )
        
        # Update joining status
        cursor.execute(
            "UPDATE employee_joining_forms SET status = 'approved' WHERE id = ?",
            (joining_id,)
        )
        
        # Also create default leave balance for the new employee
        employee_id = cursor.lastrowid
        current_year = datetime.now().year
        cursor.execute(
            "INSERT OR IGNORE INTO leave_balances (employee_id, earned_leave, special_leave, loss_of_pay, year) VALUES (?, 0, 0, 0, ?)",
            (employee_id, current_year)
        )
        
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"error": str(e)}), 500


@app.route("/download-attendance-template")
def download_attendance_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    headers = ["employee_code", "date", "status", "shift_name", "check_in", "check_out"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    sample_data = ["EMP001", "2026-04-14", "P", "Day Shift", "09:00", "18:00"]
    for col_num, value in enumerate(sample_data, 1):
        ws.cell(row=2, column=col_num, value=value)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="attendance_import_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/download-leave-balances-template")
def download_leave_balances_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "LeaveBalances"
    headers = ["employee_code", "earned_leave", "special_leave", "loss_of_pay", "year"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    # Sample: all 0s since no allot till date
    sample_data = ["EMP001", 0, 0, 0, 2026]
    for col_num, value in enumerate(sample_data, 1):
        ws.cell(row=2, column=col_num, value=value)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="leave_balances_update_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def normalize_date_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')

    text = str(value).strip()
    if not text:
        return None

    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y', '%d %b %Y', '%d %B %Y'):
        try:
            return datetime.strptime(text, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue

    return text

@app.route("/api/attendance/bulk-import", methods=["POST"])
def bulk_import_attendance():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Invalid file type. Please upload an Excel file."}), 400
    try:
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File is empty"}), 400
        headers = [str(h).strip().lower() for h in rows[0]]
        expected_headers = ["employee_code", "date", "status", "shift_name", "check_in", "check_out"]
        if headers != expected_headers:
            return jsonify({"error": "Invalid headers. Please use the attendance template."}), 400
        conn = get_connection()
        cursor = conn.cursor()
        imported = 0
        for row in rows[1:]:
            if not row[0] or not row[1] or not row[2]:
                continue
            cursor.execute("SELECT id FROM employees WHERE employee_code = ?", (str(row[0]).strip(),))
            employee = cursor.fetchone()
            if not employee:
                continue
            employee_id = employee[0]
            shift_id = None
            if row[3]:
                cursor.execute("SELECT id FROM shifts WHERE name = ?", (str(row[3]).strip(),))
                shift = cursor.fetchone()
                shift_id = shift[0] if shift else None
            date_value = normalize_date_value(row[1])
            if not date_value:
                continue
            status_value = str(row[2]).strip()
            check_in = str(row[4]).strip() if row[4] else None
            check_out = str(row[5]).strip() if row[5] else None
            effective_status = get_effective_status_for_save(employee_id, date_value, status_value, check_in, check_out, shift_id)
            cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (employee_id, date_value))
            existing = cursor.fetchone()
            if existing:
                cursor.execute(
                    "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ? WHERE id = ?",
                    (effective_status, shift_id, check_in, check_out, existing[0])
                )
            else:
                cursor.execute(
                    "INSERT INTO attendance (employee_id, date, status, shift_id, check_in, check_out) VALUES (?, ?, ?, ?, ?, ?)",
                    (employee_id, date_value, effective_status, shift_id, check_in, check_out)
                )
            imported += 1
        conn.commit()
        conn.close()
        return jsonify({"imported": imported})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/leave-balances/bulk-import", methods=["POST"])
def bulk_import_leave_balances():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Invalid file type. Please upload an Excel file."}), 400
    try:
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File is empty"}), 400
        headers = [str(h).strip().lower() for h in rows[0]]
        expected_headers = ["employee_code", "earned_leave", "special_leave", "loss_of_pay", "year"]
        if headers != expected_headers:
            return jsonify({"error": "Invalid headers. Please use the leave balances update template."}), 400
        conn = get_connection()
        cursor = conn.cursor()
        updated = 0
        current_year = datetime.now().year
        for row in rows[1:]:
            if not row[0]:
                continue
            emp_code = str(row[0]).strip()
            cursor.execute("SELECT id FROM employees WHERE employee_code = ?", (emp_code,))
            employee = cursor.fetchone()
            if not employee:
                continue
            employee_id = employee[0]
            earned = float(row[1]) if row[1] is not None else 0
            special = float(row[2]) if row[2] is not None else 0
            lop = float(row[3]) if row[3] is not None else 0
            year = int(row[4]) if row[4] is not None else current_year
            # Ensure only one row per emp+year: delete dups then insert fresh values from excel
            cursor.execute("DELETE FROM leave_balances WHERE employee_id = ? AND year = ?", (employee_id, year))
            cursor.execute(
                "INSERT INTO leave_balances (employee_id, earned_leave, special_leave, loss_of_pay, year) VALUES (?, ?, ?, ?, ?)",
                (employee_id, earned, special, lop, year)
            )
            updated += 1
        conn.commit()
        conn.close()
        return jsonify({"updated": updated})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/leave-balances")
def get_leave_balances():
    conn = get_connection()
    cursor = conn.cursor()
    current_year = datetime.now().year
    cursor.execute(
        """
        SELECT 
            e.id AS employee_id,
            e.first_name || ' ' || e.last_name AS employee_name,
            COALESCE(lb.earned_leave, 0) AS earned_leave,
            COALESCE(lb.special_leave, 0) AS special_leave,
            COALESCE(lb.loss_of_pay, 0) AS loss_of_pay,
            COALESCE(lb.year, ?) AS year
        FROM employees e
        LEFT JOIN leave_balances lb ON lb.id = (
            SELECT id FROM leave_balances lb2 
            WHERE lb2.employee_id = e.id AND lb2.year = ? 
            ORDER BY lb2.id DESC LIMIT 1
        )
        ORDER BY e.first_name, e.last_name
        """,
        (current_year, current_year)
    )
    data = dict_rows(cursor.fetchall())
    conn.close()
    return jsonify(data)

@app.route("/api/statutory")
def get_statutory():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, first_name || ' ' || last_name AS name, pf_active, esic_active, professional_tax_active AS pt_active, lwf_active FROM employees "
        "ORDER BY first_name, last_name"
    )
    data = dict_rows(cursor.fetchall())
    conn.close()
    return jsonify(data)

@app.route("/api/attendance", methods=["GET", "POST"])
def manage_attendance():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        employee_id = request.args.get('employee_id')
        start_date = request.args.get('start_date')
        end_date = request.args.get('end_date')
        query = """
            SELECT a.*, e.first_name || ' ' || e.last_name AS employee_name, s.name AS shift_name FROM attendance a 
            LEFT JOIN employees e ON a.employee_id = e.id 
            LEFT JOIN shifts s ON a.shift_id = s.id 
        """
        conditions = []
        params = []
        if employee_id:
            conditions.append("a.employee_id = ?")
            params.append(employee_id)
        fetch_start = start_date
        if start_date:
            try:
                start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
                monday = start_dt - timedelta(days=start_dt.weekday())
                fetch_start = monday.strftime('%Y-%m-%d')
            except Exception:
                pass
        if fetch_start:
            conditions.append("a.date >= ?")
            params.append(fetch_start)
        if end_date:
            conditions.append("a.date <= ?")
            params.append(end_date)
        if conditions:
            query += " WHERE " + " AND ".join(conditions)
        query += " ORDER BY a.date DESC"
        cursor.execute(query, params)
        data = dict_rows(cursor.fetchall())
        data = apply_weekend_absence_rule(data, start_date=start_date, end_date=end_date)
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    employee_id = payload.get('employee_id')
    date = payload.get('date')
    status = payload.get('status')
    shift_id = payload.get('shift_id')
    check_in = payload.get('check_in')
    check_out = payload.get('check_out')
    remarks = payload.get('remarks', '')

    required_fields = ["employee_id", "date", "status"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    effective_status = get_effective_status_for_save(employee_id, date, status, check_in, check_out, shift_id)

    # Upsert to prevent duplicate rows per employee+date (was blind INSERT before)
    cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (employee_id, date))
    existing = cursor.fetchone()
    if existing:
        cursor.execute(
            "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ?, remarks = ? WHERE id = ?",
            (effective_status, shift_id, check_in, check_out, remarks, existing[0])
        )
    else:
        cursor.execute(
            "INSERT INTO attendance (employee_id, date, status, shift_id, check_in, check_out, remarks) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (employee_id, date, effective_status, shift_id, check_in, check_out, remarks)
        )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/attendance/<int:attendance_id>", methods=["PUT"])
def update_attendance(attendance_id):
    payload = request.json or {}
    status = payload.get('status')
    shift_id = payload.get('shift_id')
    check_in = payload.get('check_in')
    check_out = payload.get('check_out')
    remarks = payload.get('remarks')

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM attendance WHERE id = ?", (attendance_id,))
    existing = cursor.fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "Attendance record not found."}), 404

    final_check_in = payload['check_in'] if 'check_in' in payload else existing['check_in']
    final_check_out = payload['check_out'] if 'check_out' in payload else existing['check_out']
    final_status = payload['status'] if 'status' in payload else existing['status']
    final_shift = payload['shift_id'] if 'shift_id' in payload else existing['shift_id']

    effective_status = get_effective_status_for_save(
        existing['employee_id'], existing['date'], final_status, final_check_in, final_check_out, final_shift
    )

    remarks_val = payload['remarks'] if 'remarks' in payload else existing['remarks']
    cursor.execute(
        "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ?, remarks = ? WHERE id = ?",
        (
            effective_status,
            final_shift,
            final_check_in,
            final_check_out,
            remarks_val,
            attendance_id
        )
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route("/api/attendance/bulk", methods=["POST"])
def bulk_mark_attendance():
    payload = request.json or {}
    date = payload.get('date')
    records = payload.get('records') or []
    if not date:
        return jsonify({"error": "date is required."}), 400
    if not isinstance(records, list) or len(records) == 0:
        return jsonify({"error": "No attendance records provided."}), 400

    conn = get_connection()
    cursor = conn.cursor()
    saved = 0
    for rec in records:
        emp_id = rec.get('employee_id')
        status = rec.get('status')
        remarks = rec.get('remarks', '')
        check_in = rec.get('check_in')
        check_out = rec.get('check_out')
        if not emp_id or not status:
            continue
        effective = get_effective_status_for_save(emp_id, date, status, check_in, check_out, rec.get('shift_id'))
        cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (emp_id, date))
        existing = cursor.fetchone()
        shift_id = rec.get('shift_id')
        if existing:
            cursor.execute(
                "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ?, remarks = ? WHERE id = ?",
                (effective, shift_id, check_in, check_out, remarks, existing[0])
            )
        else:
            cursor.execute(
                "INSERT INTO attendance (employee_id, date, status, shift_id, check_in, check_out, remarks) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (emp_id, date, effective, shift_id, check_in, check_out, remarks)
            )
        saved += 1
    conn.commit()
    conn.close()
    return jsonify({"success": True, "saved": saved})


@app.route("/api/attendance/summary")
def attendance_summary():
    conn = get_connection()
    cursor = conn.cursor()
    employee_id = request.args.get('employee_id')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    today = datetime.now().date()
    if not end_date:
        end_date = today.strftime('%Y-%m-%d')
    if not start_date:
        start_date = (today - timedelta(days=6)).strftime('%Y-%m-%d')

    start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
    end_dt = datetime.strptime(end_date, '%Y-%m-%d').date()
    if start_dt > end_dt:
        start_dt, end_dt = end_dt, start_dt

    date_labels = []
    current = start_dt
    while current <= end_dt:
        date_labels.append(current.strftime('%d %a'))
        current += timedelta(days=1)

    if employee_id:
        cursor.execute(
            "SELECT id, first_name || ' ' || last_name AS employee_name FROM employees WHERE id = ? ORDER BY first_name, last_name",
            (employee_id,)
        )
    else:
        cursor.execute(
            "SELECT id, first_name || ' ' || last_name AS employee_name FROM employees ORDER BY first_name, last_name"
        )
    employees = dict_rows(cursor.fetchall())

    fetch_start = start_date
    try:
        start_dt_obj = datetime.strptime(start_date, '%Y-%m-%d').date()
        monday = start_dt_obj - timedelta(days=start_dt_obj.weekday())
        fetch_start = monday.strftime('%Y-%m-%d')
    except Exception:
        pass

    query = "SELECT a.id, a.employee_id, a.date, a.status, a.check_in, a.check_out, a.shift_id, s.name AS shift_name FROM attendance a LEFT JOIN shifts s ON a.shift_id = s.id WHERE a.date BETWEEN ? AND ?"
    params = [fetch_start, end_date]
    if employee_id:
        query += " AND a.employee_id = ?"
        params.append(employee_id)
    cursor.execute(query, params)
    attendance_rows = dict_rows(cursor.fetchall())
    attendance_rows = apply_weekend_absence_rule(attendance_rows, start_date=start_date, end_date=end_date)

    attendance_map = {}
    for row in attendance_rows:
        if row['employee_id'] not in attendance_map:
            attendance_map[row['employee_id']] = {}
        attendance_map[row['employee_id']][row['date']] = row

    # Identify first 3 late comings per emp per full month (not just viewed range).
    # This ensures correct special "P" color even after bulk import when filter may show only part of month.
    late_allowed_set = set()
    _date_cls = date  # to avoid shadowing with later local 'date' var
    touched_months = set()
    c = start_dt
    while c <= end_dt:
        touched_months.add((c.year, c.month))
        if c.month == 12:
            c = _date_cls(c.year + 1, 1, 1)
        else:
            c = _date_cls(c.year, c.month + 1, 1)
    for emp in employees:
        eid = emp['id']
        for y, mo in touched_months:
            q = """SELECT date, check_in, check_out, shift_id FROM attendance
                   WHERE employee_id = ?
                     AND strftime('%Y', date) = ? AND strftime('%m', date) = ?"""
            cursor.execute(q, (eid, str(y), f"{mo:02d}"))
            lates = []
            for r in cursor.fetchall():
                try:
                    d = r['date']
                    cin = r['check_in']
                    cout = r['check_out']
                    sid = r['shift_id']
                except (KeyError, TypeError, IndexError):
                    d = r[0] if isinstance(r, (list, tuple)) else str(r)
                    cin = r[1] if isinstance(r, (list, tuple)) else None
                    cout = r[2] if isinstance(r, (list, tuple)) else None
                    sid = r[3] if isinstance(r, (list, tuple)) else None
                if isinstance(d, (datetime, _date_cls)):
                    d = d.strftime('%Y-%m-%d')
                if cin and is_late_arrival(cin) and is_out_on_schedule(cout, sid) and d:
                    if not isinstance(d, str):
                        d = str(d)
                    lates.append(d)
            lates.sort()
            for d in lates[:3]:
                late_allowed_set.add((eid, d))

    rows = []
    for emp in employees:
        row = {
            'id': emp['id'],
            'employee_name': emp['employee_name'],
            'present': 0,
            'absent': 0,
            'week_off': 0,
            'holiday': 0,
            'half_day': 0,
            'total': 0,
            'daily': []
        }
        daily = []
        for i in range((end_dt - start_dt).days + 1):
            current_date = start_dt + timedelta(days=i)
            date_key = current_date.strftime('%Y-%m-%d')
            record = attendance_map.get(emp['id'], {}).get(date_key)
            if record:
                status = record['status']
                row['total'] += 1
                if status == 'P':
                    row['present'] += 1
                elif status == 'A':
                    row['absent'] += 1
                elif status == 'WO':
                    row['week_off'] += 1
                elif status == 'H':
                    row['holiday'] += 1
                elif status in ('AH', 'HF'):
                    row['half_day'] += 1
                is_late_allowed = (emp['id'], date_key) in late_allowed_set
                daily.append({
                    'date': date_key,
                    'status': status,
                    'attendance_id': record['id'],
                    'is_late_allowed': is_late_allowed,
                    'check_in': record.get('check_in'),
                    'check_out': record.get('check_out'),
                    'shift_id': record.get('shift_id')
                })
            else:
                daily.append({
                    'date': date_key,
                    'status': '',
                    'attendance_id': ''
                })
        row['daily'] = daily
        rows.append(row)

    conn.close()
    return jsonify({'headers': date_labels, 'rows': rows})

@app.route("/api/leaves", methods=["GET", "POST"])
def manage_leaves():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute(
            "SELECT l.*, e.first_name || ' ' || e.last_name AS employee_name FROM leave_requests l "
            "LEFT JOIN employees e ON l.employee_id = e.id "
            "ORDER BY l.applied_on DESC"
        )
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    required_fields = ["employee_id", "leave_type", "start_date", "end_date"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    emp_id = payload.get("employee_id")
    start_d_str = payload.get("start_date")
    end_d_str = payload.get("end_date")
    ltype = payload.get("leave_type")
    reason = payload.get("reason", "")

    # Always treat create as Approved record for this admin-backend tool
    status = "Approved"

    try:
        y = datetime.strptime(start_d_str, '%Y-%m-%d').year
        ensure_leave_balance(cursor, emp_id, y)
    except Exception:
        pass

    cursor.execute(
        "INSERT INTO leave_requests (employee_id, leave_type, start_date, end_date, status, reason, applied_on) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (
            emp_id,
            ltype,
            start_d_str,
            end_d_str,
            status,
            reason,
            datetime.now().strftime("%Y-%m-%d")
        )
    )

    # Immediate effect for backend use: adjust balance + mark attendance days
    adjust_leave_balance(cursor, emp_id, ltype, start_d_str, end_d_str, reverse=False)
    mark_attendance_for_leave_days(cursor, emp_id, start_d_str, end_d_str, ltype, reason)

    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/leaves/<int:leave_id>", methods=["GET", "PUT", "DELETE"])
def manage_leave_request(leave_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM leave_requests WHERE id = ?", (leave_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    # Fetch for DELETE / PUT
    cursor.execute("SELECT * FROM leave_requests WHERE id = ?", (leave_id,))
    row = cursor.fetchone()
    existing = dict(row) if row and not isinstance(row, dict) else row

    if request.method == "DELETE":
        if existing and existing.get("status") == "Approved":
            adjust_leave_balance(
                cursor,
                existing["employee_id"],
                existing["leave_type"],
                existing["start_date"],
                existing["end_date"],
                reverse=True
            )
        cursor.execute("DELETE FROM leave_requests WHERE id = ?", (leave_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    if not existing:
        conn.close()
        return jsonify({"error": "Leave request not found."}), 404

    payload = request.json or {}
    employee_id = payload.get("employee_id", existing["employee_id"])
    leave_type = payload.get("leave_type", existing["leave_type"]).strip()
    start_date = payload.get("start_date", existing["start_date"]).strip()
    end_date = payload.get("end_date", existing["end_date"]).strip()
    status = payload.get("status", existing["status"]).strip()
    reason = payload.get("reason", existing["reason"]).strip()
    old_status = existing["status"]

    cursor.execute(
        "UPDATE leave_requests SET employee_id = ?, leave_type = ?, start_date = ?, end_date = ?, status = ?, reason = ? WHERE id = ?",
        (employee_id, leave_type, start_date, end_date, status, reason, leave_id)
    )

    # On approval, deduct from appropriate balance (Earned/Special) or record LOP
    if status == "Approved" and old_status != "Approved":
        adjust_leave_balance(cursor, employee_id, leave_type, start_date, end_date, reverse=False)
        mark_attendance_for_leave_days(cursor, employee_id, start_date, end_date, leave_type, reason)

    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/credit-earned-leaves/<int:year>/<int:month>", methods=["POST"])
def credit_earned_leaves(year, month):
    conn = get_connection()
    cursor = conn.cursor()
    
    # Get all employees
    cursor.execute("SELECT id FROM employees")
    employees = cursor.fetchall()
    
    credited_count = 0
    for emp in employees:
        emp_id = emp[0] if isinstance(emp, (list, tuple)) else emp["id"]
        
        # Check if already credited
        cursor.execute(
            "SELECT id FROM earned_leave_credits WHERE employee_id = ? AND year = ? AND month = ?",
            (emp_id, year, month)
        )
        if cursor.fetchone():
            continue  # Already credited
        
        # Probation check (3 months)
        if not is_eligible_for_earned_leave(emp_id, year, month, cursor):
            continue
        
        # Calculate present days (proxy for no major absent/leave that month)
        start_date = date(year, month, 1)
        if month == 12:
            end_date = date(year + 1, 1, 1) - timedelta(days=1)
        else:
            end_date = date(year, month + 1, 1) - timedelta(days=1)
        
        present_days = 0.0
        current = start_date
        while current <= end_date:
            cursor.execute(
                "SELECT status FROM attendance WHERE employee_id = ? AND date = ?",
                (emp_id, current.strftime('%Y-%m-%d'))
            )
            row = cursor.fetchone()
            if row:
                status = row[0] if isinstance(row, (list, tuple)) else row["status"]
                if status in ('P', 'H', 'WO'):
                    present_days += 1.0
                elif status in ('AH', 'HF'):
                    present_days += 0.5
            current += timedelta(days=1)
        
        if present_days >= 20:
            ensure_leave_balance(cursor, emp_id, year)
            # Credit 1.25
            cursor.execute(
                "INSERT INTO earned_leave_credits (employee_id, year, month, credited, credited_on) VALUES (?, ?, ?, ?, ?)",
                (emp_id, year, month, 1.25, datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            )
            # Update leave balance
            cursor.execute(
                "UPDATE leave_balances SET earned_leave = earned_leave + 1.25 WHERE employee_id = ? AND year = ?",
                (emp_id, year)
            )
            credited_count += 1
    
    conn.commit()
    conn.close()
    return jsonify({"success": True, "credited_count": credited_count})


@app.route("/api/credit-special-leaves/<int:year>", methods=["POST"])
def credit_special_leaves(year):
    """Credit 3 special leaves for the year to employees who have completed probation."""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM employees")
    employees = cursor.fetchall()
    credited_count = 0
    for emp in employees:
        emp_id = emp[0] if isinstance(emp, (list, tuple)) else emp["id"]
        # already credited this year?
        cursor.execute(
            "SELECT id FROM special_leave_credits WHERE employee_id = ? AND year = ?",
            (emp_id, year)
        )
        if cursor.fetchone():
            continue
        # check probation by end of year
        cursor.execute("SELECT joining_date FROM employees WHERE id = ?", (emp_id,))
        jrow = cursor.fetchone()
        joining = jrow[0] if jrow and isinstance(jrow, (list, tuple)) else (jrow["joining_date"] if jrow else None)
        probation_end = get_probation_end_date(joining)
        if not probation_end:
            continue
        year_end = date(year, 12, 31)
        if probation_end > year_end:
            continue
        ensure_leave_balance(cursor, emp_id, year)
        cursor.execute(
            "INSERT INTO special_leave_credits (employee_id, year, credited, credited_on) VALUES (?, ?, ?, ?)",
            (emp_id, year, 3.0, datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        )
        cursor.execute(
            "UPDATE leave_balances SET special_leave = special_leave + 3 WHERE employee_id = ? AND year = ?",
            (emp_id, year)
        )
        credited_count += 1
    conn.commit()
    conn.close()
    return jsonify({"success": True, "credited_count": credited_count})


@app.route("/api/reconcile-absences", methods=["POST"])
def reconcile_absences_with_leave():
    """Adjust leave balances based on 'A' (absent) days in attendance from 1 June 2026.
    Deduct from Earned Leave first, excess to Loss of Pay.
    Creates Approved leave_requests for the periods.
    """
    conn = get_connection()
    cursor = conn.cursor()
    as_of = '2026-06-01'
    today = datetime.now().strftime('%Y-%m-%d')

    cursor.execute("SELECT id FROM employees")
    emp_ids = [r[0] for r in cursor.fetchall()]

    adjusted = 0
    for emp_id in emp_ids:
        # current balances for 2026
        cursor.execute(
            "SELECT earned_leave, loss_of_pay FROM leave_balances WHERE employee_id = ? AND year = 2026",
            (emp_id,)
        )
        bal_row = cursor.fetchone()
        if not bal_row:
            continue
        earned = bal_row[0] or 0.0
        lop = bal_row[1] or 0.0

        # get absent dates from as_of
        cursor.execute(
            "SELECT date FROM attendance WHERE employee_id = ? AND status = 'A' AND date >= ? ORDER BY date",
            (emp_id, as_of)
        )
        dates = [r[0] for r in cursor.fetchall()]
        if not dates:
            continue

        # group consecutive periods
        periods = []
        if dates:
            s = e = dates[0]
            for d in dates[1:]:
                prev = datetime.strptime(e, '%Y-%m-%d')
                curr = datetime.strptime(d, '%Y-%m-%d')
                if (curr - prev).days == 1:
                    e = d
                else:
                    periods.append((s, e))
                    s = e = d
            periods.append((s, e))

        for pstart, pend in periods:
            d1 = datetime.strptime(pstart, '%Y-%m-%d')
            d2 = datetime.strptime(pend, '%Y-%m-%d')
            days = (d2 - d1).days + 1

            # check if this exact period already has a request (idempotent)
            cursor.execute(
                "SELECT 1 FROM leave_requests WHERE employee_id = ? AND start_date = ? AND end_date = ?",
                (emp_id, pstart, pend)
            )
            if cursor.fetchone():
                continue

            # deduct from earned
            deduct = min(days, earned)
            if deduct > 0:
                earned -= deduct
                cursor.execute(
                    """INSERT INTO leave_requests 
                    (employee_id, leave_type, start_date, end_date, status, reason, applied_on) 
                    VALUES (?, 'Earned Leave', ?, ?, 'Approved', 'Auto-adjusted from attendance (A)', ?)""",
                    (emp_id, pstart, pend, today)
                )
                adjusted += 1

            remaining = days - deduct
            if remaining > 0:
                lop += remaining
                cursor.execute(
                    """INSERT INTO leave_requests 
                    (employee_id, leave_type, start_date, end_date, status, reason, applied_on) 
                    VALUES (?, 'Loss of Pay', ?, ?, 'Approved', 'Auto-adjusted from attendance (A) - no earned leave remaining', ?)""",
                    (emp_id, pstart, pend, today)
                )
                adjusted += 1

        # update balance
        cursor.execute(
            "UPDATE leave_balances SET earned_leave = ?, loss_of_pay = ? WHERE employee_id = ? AND year = 2026",
            (earned, lop, emp_id)
        )

    conn.commit()
    conn.close()
    return jsonify({"success": True, "adjusted_periods": adjusted})


@app.route("/api/payrolls", methods=["GET", "POST"])
def manage_payrolls():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute(
            "SELECT p.*, e.first_name || ' ' || e.last_name AS employee_name FROM payrolls p "
            "LEFT JOIN employees e ON p.employee_id = e.id "
            "ORDER BY p.year DESC, p.month DESC"
        )
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    employee_id = payload.get("employee_id")
    month = payload.get("month")
    year = payload.get("year")
    if not employee_id or not month or not year:
        conn.close()
        return jsonify({"error": "Employee, month and year are required."}), 400

    cursor.execute("SELECT * FROM employees WHERE id = ?", (employee_id,))
    employee = cursor.fetchone()
    if not employee:
        conn.close()
        return jsonify({"error": "Employee not found."}), 404

    ctc = float(employee["ctc"] or 0)
    pf_active = bool(employee["pf_active"])
    esic_active = bool(employee["esic_active"])
    professional_tax_active = bool(employee["professional_tax_active"])
    lwf_active = bool(employee["lwf_active"])
    payroll = calculate_pay_components(ctc, pf_active, esic_active, professional_tax_active, lwf_active)
    cursor.execute(
        "SELECT id FROM payrolls WHERE employee_id = ? AND month = ? AND year = ?",
        (employee_id, month, year)
    )
    if cursor.fetchone():
        conn.close()
        return jsonify({"error": "Payroll already generated for this employee and period."}), 400

    cursor.execute(
        "INSERT INTO payrolls (employee_id, month, year, ctc, basic, da, hra, conveyance, special_allowance, gross, pf, esic, professional_tax, lwf, total_deductions, net_pay) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            employee_id,
            month,
            year,
            ctc,
            payroll["basic"],
            payroll["da"],
            payroll["hra"],
            payroll["conveyance"],
            payroll["special_allowance"],
            payroll["gross"],
            payroll["pf"],
            payroll["esic"],
            payroll["professional_tax"],
            payroll["lwf"],
            payroll["total_deductions"],
            payroll["net_pay"]
        )
    )
    conn.commit()
    payroll_id = cursor.lastrowid
    cursor.execute("SELECT * FROM payrolls WHERE id = ?", (payroll_id,))
    payroll_record = dict(cursor.fetchone())
    conn.close()
    payroll_record["employee_name"] = f"{employee['first_name']} {employee['last_name']}"
    return jsonify(payroll_record)


# ============================================================================
# ================  Bank Bulk Salary File Generator — routes  ==============
# ============================================================================

@app.route("/api/bank-employees", methods=["GET", "POST"])
def bank_employees():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        company = request.args.get("company")
        if company:
            cursor.execute("SELECT * FROM bank_employees WHERE company = ? ORDER BY name", (company,))
        else:
            cursor.execute("SELECT * FROM bank_employees ORDER BY company, name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = (payload.get("name") or "").strip()
    ifsc = (payload.get("ifsc") or "").strip().upper()
    account = str(payload.get("account") or "").strip()
    company = payload.get("company") or ""
    if not name or not ifsc or not account or company not in BANK_COMPANIES:
        conn.close()
        return jsonify({"error": "Name, IFSC, Account No. and a valid Company are all required."}), 400
    try:
        cursor.execute(
            "INSERT INTO bank_employees (name, ifsc, account, company) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(name, company) DO UPDATE SET ifsc=excluded.ifsc, account=excluded.account",
            (name, ifsc, account, company)
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": f"'{name}' already exists under '{company}'."}), 400
    conn.close()
    return jsonify({"success": True})


@app.route("/api/bank-employees/<int:emp_id>", methods=["PUT", "DELETE"])
def bank_employee_detail(emp_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "DELETE":
        cursor.execute("DELETE FROM bank_employees WHERE id = ?", (emp_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    payload = request.json or {}
    name = (payload.get("name") or "").strip()
    ifsc = (payload.get("ifsc") or "").strip().upper()
    account = str(payload.get("account") or "").strip()
    company = payload.get("company") or ""
    if not name or not ifsc or not account or company not in BANK_COMPANIES:
        conn.close()
        return jsonify({"error": "Name, IFSC, Account No. and a valid Company are all required."}), 400
    try:
        cursor.execute(
            "UPDATE bank_employees SET name=?, ifsc=?, account=?, company=? WHERE id=?",
            (name, ifsc, account, company, emp_id)
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": f"'{name}' already exists under '{company}'."}), 400
    conn.close()
    return jsonify({"success": True})


@app.route("/api/bank-employees/bulk-upload", methods=["POST"])
def bank_employees_bulk_upload():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    company = request.form.get("company") or ""
    if company not in BANK_COMPANIES:
        return jsonify({"error": "Please select a company to assign these employees to."}), 400

    try:
        employees, detected = read_bank_employees_from_template(file)
    except Exception as exc:
        return jsonify({"error": f"Could not read that file: {exc}"}), 400

    if not employees:
        return jsonify({"error": "No employee rows were found in that file."}), 400

    if detected and detected != company:
        company = detected  # the file's own company code wins, mirrors original tool behaviour

    conn = get_connection()
    cursor = conn.cursor()
    count = 0
    for name, ifsc, account in employees:
        cursor.execute(
            "INSERT INTO bank_employees (name, ifsc, account, company) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(name, company) DO UPDATE SET ifsc=excluded.ifsc, account=excluded.account",
            (name, ifsc.upper(), account, company)
        )
        count += 1
    conn.commit()
    conn.close()
    return jsonify({"count": count, "company": company})


@app.route("/api/bank-remarks")
def bank_remarks():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT remark FROM bank_remarks ORDER BY use_count DESC, last_used DESC")
    remarks = [r[0] for r in cursor.fetchall()]
    conn.close()
    return jsonify(remarks)


@app.route("/api/bank-batches", methods=["GET", "POST"])
def bank_batches():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute(
            "SELECT b.id, b.batch_date, b.company, b.created_at, b.filename, "
            "COUNT(t.id) AS entry_count, COALESCE(SUM(t.amount), 0) AS total_amount "
            "FROM bank_batches b LEFT JOIN bank_transactions t ON t.batch_id = b.id "
            "GROUP BY b.id ORDER BY b.id DESC"
        )
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    company = payload.get("company") or ""
    date_str = (payload.get("date") or "").strip()
    entries = payload.get("entries") or []
    if company not in BANK_COMPANIES:
        conn.close()
        return jsonify({"error": "Please select a valid company."}), 400
    if not BANK_DATE_RE.match(date_str):
        conn.close()
        return jsonify({"error": "Date must be exactly in DD/MM/YYYY format."}), 400
    try:
        datetime.strptime(date_str, "%d/%m/%Y")
    except ValueError:
        conn.close()
        return jsonify({"error": "That is not a real calendar date."}), 400
    if not entries:
        conn.close()
        return jsonify({"error": "Add at least one employee entry first."}), 400

    cfg = BANK_COMPANY_CONFIG[company]
    filename = f"{cfg['code']}_{date_str.replace('/', '-')}.xlsx"

    cursor.execute(
        "INSERT INTO bank_batches (batch_date, company, created_at, filename) VALUES (?, ?, ?, ?)",
        (date_str, company, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), filename)
    )
    batch_id = cursor.lastrowid
    for e in entries:
        mode = bank_payment_mode_for_ifsc(e.get("ifsc", ""))
        remarks = (e.get("remarks") or BANK_DEFAULT_REMARK).strip() or BANK_DEFAULT_REMARK
        cursor.execute(
            "INSERT INTO bank_transactions (batch_id, name, ifsc, account, amount, mode, remarks) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (batch_id, e.get("name"), e.get("ifsc"), e.get("account"), float(e.get("amount", 0)), mode, remarks)
        )
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute(
            "INSERT INTO bank_remarks (remark, last_used, use_count) VALUES (?, ?, 1) "
            "ON CONFLICT(remark) DO UPDATE SET last_used=excluded.last_used, use_count=use_count+1",
            (remarks, now)
        )
    conn.commit()
    conn.close()
    return jsonify({"id": batch_id, "filename": filename})


@app.route("/api/bank-batches/<int:batch_id>/download")
def bank_batch_download(batch_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT batch_date, company, filename FROM bank_batches WHERE id = ?", (batch_id,))
    info = cursor.fetchone()
    if not info:
        conn.close()
        return jsonify({"error": "Batch not found."}), 404
    cursor.execute(
        "SELECT name, ifsc, account, amount, mode, remarks FROM bank_transactions WHERE batch_id = ? ORDER BY id",
        (batch_id,)
    )
    rows = cursor.fetchall()
    conn.close()

    entries = [{"name": r["name"], "ifsc": r["ifsc"], "account": r["account"],
                "amount": r["amount"], "remarks": r["remarks"]} for r in rows]
    output = generate_bank_excel(info["company"], info["batch_date"], entries)
    return send_file(
        output, as_attachment=True,
        download_name=info["filename"] or f"bank_batch_{batch_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ============================================================================
# ================  Payroll Processing System — routes  =====================
# ============================================================================

@app.route("/api/payroll-master-employees", methods=["GET", "POST"])
def payroll_master_employees():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM payroll_master_employees ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = (payload.get("name") or "").strip().upper()
    code = (payload.get("code") or "").strip().upper()
    office = (payload.get("office") or "").strip()
    try:
        ctc = float(payload.get("ctc"))
        if ctc <= 0:
            raise ValueError
    except (TypeError, ValueError):
        conn.close()
        return jsonify({"error": "Enter a valid positive CTC amount."}), 400
    if not name:
        conn.close()
        return jsonify({"error": "Employee name is required."}), 400
    if not office:
        conn.close()
        return jsonify({"error": "Please select the employee's office."}), 400
    try:
        cursor.execute(
            "INSERT INTO payroll_master_employees (code, name, office, ctc) VALUES (?, ?, ?, ?)",
            (code, name, office, ctc)
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": f"'{name}' already exists."}), 400
    conn.close()
    return jsonify({"success": True})


@app.route("/api/payroll-master-employees/<int:emp_id>", methods=["PUT", "DELETE"])
def payroll_master_employee_detail(emp_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "DELETE":
        cursor.execute("DELETE FROM payroll_master_employees WHERE id = ?", (emp_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    payload = request.json or {}
    name = (payload.get("name") or "").strip().upper()
    code = (payload.get("code") or "").strip().upper()
    office = (payload.get("office") or "").strip()
    try:
        ctc = float(payload.get("ctc"))
        if ctc <= 0:
            raise ValueError
    except (TypeError, ValueError):
        conn.close()
        return jsonify({"error": "Enter a valid positive CTC amount."}), 400
    if not name:
        conn.close()
        return jsonify({"error": "Employee name is required."}), 400
    if not office:
        conn.close()
        return jsonify({"error": "Please select the employee's office."}), 400
    try:
        cursor.execute(
            "UPDATE payroll_master_employees SET code=?, name=?, office=?, ctc=? WHERE id=?",
            (code, name, office, ctc, emp_id)
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": f"'{name}' already exists."}), 400
    conn.close()
    return jsonify({"success": True})


@app.route("/api/payroll-master/run", methods=["POST"])
def payroll_master_run():
    payload = request.json or {}
    entries = payload.get("entries") or []
    if not entries:
        return jsonify({"error": "Add at least one employee before running payroll."}), 400

    conn = get_connection()
    cursor = conn.cursor()
    results = []
    for entry in entries:
        emp_id = entry.get("employee_id")
        cursor.execute(
            "SELECT e.employee_code AS code, e.first_name, e.last_name, e.ctc, o.name AS office_name "
            "FROM employees e LEFT JOIN offices o ON e.office_id = o.id WHERE e.id = ?",
            (emp_id,)
        )
        emp = cursor.fetchone()
        if not emp:
            conn.close()
            return jsonify({"error": "One of the selected employees no longer exists."}), 400
        name = f"{emp['first_name']} {emp['last_name']}".strip()
        if emp["ctc"] is None:
            conn.close()
            return jsonify({"error": f"{name} has no CTC set in Employee Records."}), 400
        try:
            month_days = float(entry.get("month_days"))
            present_days = float(entry.get("present_days"))
            if month_days <= 0 or present_days < 0 or present_days > month_days:
                raise ValueError
        except (TypeError, ValueError):
            conn.close()
            return jsonify({"error": f"Invalid Month Days / Present Days for {name}."}), 400
        result = calculate_master_payroll(
            name, emp["ctc"], month_days, present_days, emp["code"] or "", emp["office_name"] or ""
        )
        results.append(result)
    conn.close()
    return jsonify(results)


@app.route("/api/payroll-master/save", methods=["POST"])
def payroll_master_save():
    payload = request.json or {}
    month = payload.get("month")
    year = str(payload.get("year") or "").strip()
    records = payload.get("records") or []
    if month not in PAYROLL_MONTH_NAMES:
        return jsonify({"error": "Please select a valid month."}), 400
    if not year.isdigit():
        return jsonify({"error": "Please enter a valid year."}), 400
    if not records:
        return jsonify({"error": "Run the payroll calculation first."}), 400

    conn = get_connection()
    cursor = conn.cursor()

    batch_no = next_payroll_batch_no(cursor, month, year)
    cursor.execute(
        "INSERT INTO payroll_runs (month, year, batch_no, saved_on) VALUES (?, ?, ?, ?)",
        (month, year, batch_no, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    )
    run_id = cursor.lastrowid
    for r in records:
        cursor.execute(
            "INSERT INTO payroll_run_records (run_id, name, code, office, ctc, month_days, present_days, "
            "gross, basic_da, hra, other_allowance, total_earning, employee_pf, employee_esic, pt, "
            "total_deduction, net_pay, employer_pf, employer_esic, lwf, cost_to_company, additional_ctc) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                run_id, r.get("Name"), r.get("Employee Code"), r.get("Office"), r.get("CTC"),
                r.get("Month Days"), r.get("Present Days"), r.get("Gross"), r.get("Basic+DA"),
                r.get("HRA"), r.get("Other Allowance"), r.get("Total Earning"), r.get("Employee PF"),
                r.get("Employee ESIC"), r.get("PT"), r.get("Total Deduction"), r.get("Net Pay"),
                r.get("Employer PF"), r.get("Employer ESIC"), r.get("LWF"), r.get("Cost To Company"),
                r.get("Additional CTC")
            )
        )
    conn.commit()
    conn.close()
    return jsonify({"success": True, "month": month, "year": year, "batch_no": batch_no})


@app.route("/api/payroll-batches")
def payroll_batches():
    month_f = request.args.get("month")
    year_f = request.args.get("year")
    conn = get_connection()
    cursor = conn.cursor()
    query = (
        "SELECT run.id, run.month, run.year, run.batch_no, run.saved_on, COUNT(rec.id) AS record_count "
        "FROM payroll_runs run LEFT JOIN payroll_run_records rec ON rec.run_id = run.id WHERE 1=1"
    )
    params = []
    if month_f:
        query += " AND run.month = ?"
        params.append(month_f)
    if year_f:
        query += " AND run.year = ?"
        params.append(year_f)
    query += " GROUP BY run.id ORDER BY run.id DESC"
    cursor.execute(query, params)
    data = dict_rows(cursor.fetchall())
    conn.close()
    return jsonify(data)


@app.route("/api/payroll-batches/<int:run_id>/records")
def payroll_batch_records(run_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT rec.*, run.month AS month, run.year AS year, run.batch_no AS batch_no "
        "FROM payroll_run_records rec JOIN payroll_runs run ON rec.run_id = run.id "
        "WHERE run.id = ? ORDER BY rec.name",
        (run_id,)
    )
    rows = cursor.fetchall()
    conn.close()
    if not rows:
        return jsonify({"error": "Batch not found or has no records."}), 404
    return jsonify([payroll_record_row_to_dict(row) for row in rows])


@app.route("/api/payroll-master/records")
def payroll_master_records():
    name_f = request.args.get("name")
    month_f = request.args.get("month")
    year_f = request.args.get("year")
    office_f = request.args.get("office")

    conn = get_connection()
    cursor = conn.cursor()
    query = (
        "SELECT rec.*, run.month AS month, run.year AS year, run.batch_no AS batch_no, run.id AS run_id "
        "FROM payroll_run_records rec JOIN payroll_runs run ON rec.run_id = run.id WHERE 1=1"
    )
    params = []
    if name_f:
        query += " AND rec.name = ?"
        params.append(name_f)
    if month_f:
        query += " AND run.month = ?"
        params.append(month_f)
    if year_f:
        query += " AND run.year = ?"
        params.append(year_f)
    if office_f:
        query += " AND rec.office = ?"
        params.append(office_f)
    query += " ORDER BY run.year DESC, run.month DESC, rec.name"
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()

    flat = [payroll_record_row_to_dict(row) for row in rows]
    return jsonify(flat)


@app.route("/api/payroll-master/records/delete", methods=["POST"])
def payroll_master_records_delete():
    payload = request.json or {}
    ids = payload.get("ids") or []
    if not ids:
        return jsonify({"error": "No records selected."}), 400
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(f"SELECT DISTINCT run_id FROM payroll_run_records WHERE id IN ({','.join('?' * len(ids))})", ids)
    run_ids = [r["run_id"] for r in cursor.fetchall()]
    cursor.execute(f"DELETE FROM payroll_run_records WHERE id IN ({','.join('?' * len(ids))})", ids)
    # Clean up any runs left with no records
    for run_id in run_ids:
        cursor.execute("SELECT COUNT(*) AS c FROM payroll_run_records WHERE run_id = ?", (run_id,))
        if cursor.fetchone()["c"] == 0:
            cursor.execute("DELETE FROM payroll_runs WHERE id = ?", (run_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True, "deleted": len(ids)})


@app.route("/api/payroll-master/records/delete-month", methods=["POST"])
def payroll_master_records_delete_month():
    payload = request.json or {}
    month = payload.get("month")
    year = str(payload.get("year") or "").strip()
    if not month or not year:
        return jsonify({"error": "Select a month and year to delete."}), 400
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM payroll_runs WHERE month = ? AND year = ?", (month, year))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route("/api/payroll-master/export", methods=["POST"])
def payroll_master_export():
    payload = request.json or {}
    records = payload.get("records") or []
    include_period = bool(payload.get("include_period"))
    if not records:
        return jsonify({"error": "No records to export."}), 400
    columns = (["Batch No", "Month", "Year"] if include_period else []) + PAYROLL_RESULT_COLUMNS
    output = write_payroll_excel(columns, records)
    filename = f"Payroll_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/generate-offer", methods=["POST"])
def generate_offer():
    payload = request.json or {}
    required = ["offer_date", "prefix", "name", "position", "joining_date", "location", "department", "monthly_ctc"]
    for field in required:
        if not payload.get(field):
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    ensure_template()
    monthly_ctc = float(payload.get("monthly_ctc"))
    annual_ctc = monthly_ctc * 12
    replacements = {
        "[DATE]": payload.get("offer_date"),
        "[PREFIX]": payload.get("prefix"),
        "[NAME]": payload.get("name"),
        "[POSITION]": payload.get("position"),
        "[JOINING_DATE]": payload.get("joining_date"),
        "[LOCATION]": payload.get("location"),
        "[DEPARTMENT]": payload.get("department"),
        "[CTC_A]": str(int(annual_ctc)),
        "[CTC_A_WORD]": num2words(int(annual_ctc), lang="en_IN").title() + " Only",
        "[CTC_M]": str(int(monthly_ctc)),
        "[CTC_M_WORD]": num2words(int(monthly_ctc), lang="en_IN").title() + " Only"
    }
    bold_keys = ["[NAME]", "[POSITION]", "[JOINING_DATE]", "[LOCATION]", "[DEPARTMENT]", "[CTC_A]", "[CTC_A_WORD]", "[CTC_M]", "[CTC_M_WORD]"]
    temp_path = os.path.join(BASE_DIR, f"offer_letter_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx")
    shutil.copy(TEMPLATE_PATH, temp_path)
    doc = Document(temp_path)
    replace_placeholders(doc, replacements, bold_keys)
    doc.save(temp_path)
    with open(temp_path, 'rb') as f:
        file_bytes = f.read()
    os.remove(temp_path)
    return send_file(
        io.BytesIO(file_bytes),
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=f"{payload.get('name')}_Offer_Letter_{payload.get('position')}.docx"
    )

init_db()
ensure_template()

if __name__ == '__main__':
    debug_mode = os.environ.get('FLASK_DEBUG', 'false').lower() == 'true'
    app.run(debug=debug_mode, port=int(os.environ.get('PORT', 5000)))
